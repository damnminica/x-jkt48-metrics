"""Excel gabungan banyak member dari DATA_DIR/<HANDLE>_v2.json.

Sheet: Ringkasan (ranking per member), Per Tim, Bulanan, Interaksi Member, Top 50, Data.
Ditulis dengan mode write-only openpyxl (baris demi baris) supaya muat di VPS 1 GB.
Dipanggil worker.py di akhir batch; bisa juga manual:
    python3 export_combined.py DATA_DIR OUT.xlsx [CUTOFF_YYYY-MM-DD] [HANDLE,HANDLE,...]
"""
import csv
import json
import sys
from collections import Counter
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE = Path(__file__).parent
WIB = timedelta(hours=7)
HARI = ["Senin", "Selasa", "Rabu", "Kamis", "Jumat", "Sabtu", "Minggu"]
UTAMA, REPLY, REPOST = "Post utama", "Reply", "Repost"
TEAM_ORDER = ["Dream", "Love", "Passion", "Trainee"]

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
SUB_FILL = PatternFill("solid", fgColor="D9E1F2")
BOX = Border(bottom=Side(style="thin", color="BFBFBF"))
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
NUM, PCT = "#,##0", "0.00%"
GREY, BLUE, GREEN, LINK = "595959", "0000FF", "008000", "0563C1"
_FONTS = {}


def f(bold=False, color="000000", size=10, italic=False):
    key = (bold, color, size, italic)
    if key not in _FONTS:  # objek Font dipakai ulang; jutaan sel tidak membuat jutaan objek
        _FONTS[key] = Font(name=FONT, bold=bold, color=color, size=size, italic=italic)
    return _FONTS[key]


def clean_text(s):
    return ILLEGAL_CHARACTERS_RE.sub("", s.replace("\n", " "))


class Sheet:
    """Pembungkus sheet write-only: baris ditambahkan berurutan, nomor baris dilacak."""

    def __init__(self, wb, title, col_widths):
        self.ws = wb.create_sheet(title)
        for i, w in enumerate(col_widths, 1):
            self.ws.column_dimensions[get_column_letter(i)].width = w
        self.row = 0

    def c(self, value, fmt=None, bold=False, color="000000", italic=False, size=10, fill=None, box=False,
          text=False, align=None):
        cell = WriteOnlyCell(self.ws, value=clean_text(value) if text and isinstance(value, str) else value)
        if text:
            cell.data_type = "s"  # teks yang diawali '=' tetap teks, bukan rumus
        cell.font = f(bold, color, size, italic)
        if fmt:
            cell.number_format = fmt
        if fill:
            cell.fill = fill
        if box:
            cell.border = BOX
        if align:
            cell.alignment = align
        return cell

    def append(self, cells=()):
        self.ws.append(list(cells))
        self.row += 1
        return self.row

    def header(self, labels, offset=0):
        self.ws.row_dimensions[self.row + 1].height = 32
        return self.append([None] * offset + [self.c(h, bold=True, color="FFFFFF", fill=HEAD_FILL, align=CENTER)
                                              for h in labels])

    def title(self, text, sub=None):
        self.append([self.c(text, bold=True, size=12)])
        self.append([self.c(sub, italic=True, color=GREY)] if sub else [])
        self.append()

    def table_row(self, values, fmts=None, colors=None, total=False, offset=0):
        """values: list nilai/rumus; fmts/colors: dict kolom(1-based)->format/warna."""
        cells = [None] * offset
        for j, v in enumerate(values, 1):
            cells.append(self.c(v, fmt=(fmts or {}).get(j, NUM), bold=total, color=(colors or {}).get(j, "000000"),
                                fill=SUB_FILL if total else None, box=True,
                                text=isinstance(v, str) and not v.startswith("=")))
        return self.append(cells)


def group_of(kind):
    return UTAMA if kind in ("Original", "Quote") else REPOST if kind == "Repost" else REPLY


def load(members, data_dir, cutoff):
    """Gabungkan tweet semua member (>= cutoff, WIB); simpan hanya field yang dipakai."""
    rows, profiles, missing = [], {}, []
    for m in members:
        path = Path(data_dir) / f"{m['handle_x']}_v2.json"
        if not path.exists():
            missing.append(m["handle_x"])
            continue
        raw = json.loads(path.read_text())
        profiles[m["handle_x"]] = raw.get("profile") or {}
        for t in raw["tweets"].values():
            utc = datetime.fromisoformat(t["created_at"]).replace(tzinfo=None)
            if utc + WIB >= cutoff:  # cutoff = tanggal WIB
                rows.append({"wib": utc + WIB, "member": m, "id": t["id"], "kind": t["kind"],
                             "reply_to": t["reply_to"], "orig_author": t["orig_author"], "orig_id": t["orig_id"],
                             "metrics": t["metrics"], "text": t["text"]})
        del raw
    rows.sort(key=lambda r: (r["member"]["handle_x"].lower(), r["wib"]))
    return rows, profiles, missing


def _id_num(x):
    return f"{x:,.0f}".replace(",", ".")


def summary_text(members, data_dir, cutoff, top=10):
    """Ringkasan teks (dihitung di Python) untuk Telegram, karena pratinjau Excel di HP tidak menghitung rumus."""
    rows, _, missing = load(members, data_dir, cutoff)
    stats = {}
    for r in rows:
        if group_of(r["kind"]) == UTAMA:
            s = stats.setdefault(r["member"]["handle_x"], {"m": r["member"], "n": 0, "views": 0, "likes": 0})
            s["n"] += 1
            s["views"] += r["metrics"]["views"]
            s["likes"] += r["metrics"]["likes"]
    ranked = sorted(stats.values(), key=lambda s: -s["views"] / s["n"])
    lines = [f"🏆 Rata-rata views per post utama (sejak {cutoff:%d %b %Y} WIB), {len(ranked)} member:"]
    for i, s in enumerate(ranked[:top], 1):
        lines.append(f"{i}. {s['m']['panggilan']} — {_id_num(s['views'] / s['n'])} views, "
                     f"{_id_num(s['likes'] / s['n'])} like/post ({s['n']} post)")
    total_n = sum(s["n"] for s in ranked)
    total_v = sum(s["views"] for s in ranked)
    lines.append(f"Total: {_id_num(total_n)} post utama, {_id_num(total_v)} views.")
    if missing:
        lines.append(f"Tanpa data: {', '.join(missing)}")
    return "\n".join(lines)


def build_workbook(members, data_dir, cutoff, out):
    rows, profiles, missing = load(members, data_dir, cutoff)
    members = [m for m in members if m["handle_x"] in profiles]
    if not members:
        raise ValueError("Tidak ada data member untuk diekspor")
    last_date = max(r["wib"] for r in rows) if rows else cutoff
    handle_case = {m["handle_x"].lower(): m["handle_x"] for m in members}
    ordered = sorted(members, key=lambda m: -(profiles[m["handle_x"]].get("followers_count") or 0))
    months = sorted({r["wib"].strftime("%Y-%m") for r in rows})
    n, n_months = len(members), max(1, len(months))
    last = len(rows) + 1  # baris terakhir sheet Data

    def rng(col):
        return f"Data!${col}$2:${col}${last}"

    def crit_member(cell, group=None):
        return f"{rng('A')},{cell}" + (f',{rng("I")},"{group}"' if group else "")

    wb = Workbook(write_only=True)

    # ── Ringkasan: satu baris per member ──────────────────────────────
    s = Sheet(wb, "Ringkasan", [7, 11, 9, 17, 11, 9, 8, 8, 14, 12, 11, 10, 11, 8, 10])
    s.ws.freeze_panes = "E5"
    s.title("Metrik X member JKT48",
            f"Periode {cutoff:%d %b %Y} – {last_date:%d %b %Y} (WIB) · {n} member · post utama = Original + Quote; "
            "repost tidak dihitung di metrik")
    s.header(["Rank views", "Panggilan", "Tim", "Handle", "Follower", "Post utama", "Reply", "Repost",
              "Total views (post utama)", "Rata-rata views", "Rata-rata like", "Rata-rata RT", "Total bookmark",
              "ER", "Post utama / bulan"])
    r1, rlast = 5, 4 + n
    fmts = {14: PCT, 15: "0.0"}
    for i, m in zip(range(r1, rlast + 1), ordered):
        h, cu = f"$D{i}", crit_member(f"$D{i}", UTAMA)
        s.table_row([
            f'=COUNTIF($J${r1}:$J${rlast},">"&J{i})+1',  # setara RANK(), jalan juga di Google Sheets/LibreOffice
            m["panggilan"], m["tim"], m["handle_x"], profiles[m["handle_x"]].get("followers_count") or 0,
            f"=COUNTIFS({cu})", f"=COUNTIFS({crit_member(h, REPLY)})", f"=COUNTIFS({crit_member(h, REPOST)})",
            f"=SUMIFS({rng('O')},{cu})", f"=IF(F{i}>0,I{i}/F{i},0)",
            f"=IF(F{i}>0,SUMIFS({rng('K')},{cu})/F{i},0)", f"=IF(F{i}>0,SUMIFS({rng('L')},{cu})/F{i},0)",
            f"=SUMIFS({rng('P')},{cu})", f"=IF(I{i}>0,SUMIFS({rng('Q')},{cu})/I{i},0)", f"=F{i}/{n_months}",
        ], fmts=fmts, colors={5: BLUE})
    tot = rlast + 1
    all_main = f'{rng("I")},"{UTAMA}"'
    s.table_row([None, "Total / rata-rata", None, None, f"=SUM(E{r1}:E{rlast})", f"=SUM(F{r1}:F{rlast})",
                 f"=SUM(G{r1}:G{rlast})", f"=SUM(H{r1}:H{rlast})", f"=SUM(I{r1}:I{rlast})",
                 f"=IF(F{tot}>0,I{tot}/F{tot},0)", f"=IF(F{tot}>0,SUMIFS({rng('K')},{all_main})/F{tot},0)",
                 f"=IF(F{tot}>0,SUMIFS({rng('L')},{all_main})/F{tot},0)", f"=SUM(M{r1}:M{rlast})",
                 f"=IF(I{tot}>0,SUMIFS({rng('Q')},{all_main})/I{tot},0)", f"=F{tot}/{n_months}"],
                fmts=fmts, total=True)
    s.ws.auto_filter.ref = f"A4:O{rlast}"
    s.append()
    notes = [
        "• Rank views = peringkat rata-rata views per post utama (1 = tertinggi). Pakai filter header untuk mengurutkan.",
        "• Follower (biru) = angka saat diambil; X tidak menyediakan riwayat follower.",
        "• Like, views, dll. adalah angka saat diambil, bukan saat tweet diposting. ER = (like+RT+reply+quote) ÷ views.",
        f"• Post utama / bulan dibagi {n_months} bulan kalender yang muncul di data (bulan pertama & terakhir tidak penuh).",
        "• Data dari search 'from:akun' per minggu; search X bisa sedikit tidak lengkap.",
        "• Semua angka adalah rumus yang dihitung saat file dibuka di Excel/Google Sheets; "
        "pratinjau di HP bisa tampil kosong/0.",
    ]
    if missing:
        notes.append(f"• Tidak ada data (belum ditarik): {', '.join(missing)}")
    s.append([s.c("Catatan", bold=True)])
    for note in notes:
        s.append([s.c(note, italic=True, color=GREY, text=True)])
    ch = BarChart()
    ch.type, ch.title, ch.height, ch.width, ch.legend = "bar", "Rata-rata views per post utama", 14, 18, None
    ch.add_data(Reference(s.ws, min_col=10, min_row=4, max_row=rlast), titles_from_data=True)
    ch.set_categories(Reference(s.ws, min_col=2, min_row=r1, max_row=rlast))
    ch.x_axis.scaling.orientation = "maxMin"
    s.ws.add_chart(ch, "Q4")

    # ── Per Tim ───────────────────────────────────────────────────────
    t = Sheet(wb, "Per Tim", [10, 10, 11, 9, 15, 14, 13, 8, 12])
    t.title("Perbandingan per tim", "Dihitung dari sheet Data; post utama = Original + Quote.")
    t.header(["Tim", "Jumlah member", "Post utama", "Reply", "Total views", "Rata-rata views / post",
              "Rata-rata like / post", "ER", "Post utama / member"])
    teams = sorted({m["tim"] for m in members}, key=lambda x: TEAM_ORDER.index(x) if x in TEAM_ORDER else 9)
    for i, team in enumerate(teams, 5):
        ct = f'{rng("C")},$A{i}'
        t.table_row([team, f"=COUNTIFS(Ringkasan!$C${r1}:$C${rlast},$A{i})",
                     f'=COUNTIFS({ct},{all_main})', f'=COUNTIFS({ct},{rng("I")},"{REPLY}")',
                     f'=SUMIFS({rng("O")},{ct},{all_main})', f"=IF(C{i}>0,E{i}/C{i},0)",
                     f'=IF(C{i}>0,SUMIFS({rng("K")},{ct},{all_main})/C{i},0)',
                     f'=IF(E{i}>0,SUMIFS({rng("Q")},{ct},{all_main})/E{i},0)', f"=IF(B{i}>0,C{i}/B{i},0)"],
                    fmts={8: PCT})

    # ── Bulanan: matriks member x bulan ───────────────────────────────
    b = Sheet(wb, "Bulanan", [11, 17] + [10] * len(months))
    b.ws.freeze_panes = "C6"
    b.title("Tren bulanan per member (post utama)",
            "Atas: rata-rata views per post. Bawah: jumlah post utama. Bulan pertama & terakhir tidak penuh.")
    for label, kind in [("Rata-rata views per post utama", "views"), ("Jumlah post utama", "count")]:
        b.append([b.c(label, bold=True)])
        hdr = b.header(["Panggilan", "Handle"] + months)
        for m in ordered:
            i = b.row + 1
            vals = [m["panggilan"], m["handle_x"]]
            for j in range(3, 3 + len(months)):
                crit = f'{rng("A")},$B{i},{rng("E")},{get_column_letter(j)}${hdr},{all_main}'
                vals.append(f"=COUNTIFS({crit})" if kind == "count"
                            else f"=IFERROR(SUMIFS({rng('O')},{crit})/COUNTIFS({crit}),0)")
            b.table_row(vals)
        b.append()
        b.append()

    # ── Interaksi Member: siapa membalas siapa ────────────────────────
    it = Sheet(wb, "Interaksi Member", [4, 17, 17, 12, 15])
    it.title("Siapa membalas siapa (antar member)",
             "Reply dari member ke akun member lain; urutan 40 pasangan teratas dibuat saat ekspor.")
    it.header(["#", "Dari", "Membalas", "Jumlah reply", "Rata-rata views reply"])
    pairs = Counter((r["member"]["handle_x"], handle_case[r["reply_to"].lower()]) for r in rows
                    if r["kind"] == "Reply" and r["reply_to"].lower() in handle_case)
    for k, ((src, dst), _) in enumerate(pairs.most_common(40), 1):
        i = it.row + 1
        crit = f'{rng("A")},$B{i},{rng("J")},$C{i},{rng("H")},"Reply"'
        it.table_row([k, src, dst, f"=COUNTIFS({crit})", f"=IF(D{i}>0,SUMIFS({rng('O')},{crit})/D{i},0)"])
    if not pairs:
        it.append([it.c("Belum ada reply antar member di data ini.", italic=True, color=GREY)])

    # ── Top 50 post utama ─────────────────────────────────────────────
    tp = Sheet(wb, "Top 50", [4, 11, 16, 9, 11, 9, 9, 8, 10, 50, 12])
    tp.title("50 post utama dengan views tertinggi (semua member)",
             "Urutan dibuat saat ekspor; angka hijau = tautan ke sheet Data.")
    tp.header(["#", "Panggilan", "Tanggal (WIB)", "Kategori", "Views", "Like", "Retweet", "Reply", "Bookmark",
               "Teks", "Link"])
    main_rows = [(i, r) for i, r in enumerate(rows, 2) if group_of(r["kind"]) == UTAMA]
    for k, (src, _) in enumerate(sorted(main_rows, key=lambda x: -x[1]["metrics"]["views"])[:50], 1):
        vals = [k] + [f"=Data!{c}{src}" for c in "BDHOKLMPR"] + [f'=HYPERLINK(Data!S{src},"buka")']
        tp.table_row(vals, fmts={3: "yyyy-mm-dd hh:mm", 4: "@", 10: "@", 11: "@"},
                     colors={j: GREEN for j in range(2, 11)} | {11: LINK})

    # ── Data (terakhir; baris di-stream) ──────────────────────────────
    d = Sheet(wb, "Data", [16, 10, 9, 16, 9, 8, 5, 18, 10, 18, 9, 9, 8, 7, 11, 9, 10, 60, 45])
    d.ws.freeze_panes = "B2"
    d.header(["Handle", "Panggilan", "Tim", "Tanggal (WIB)", "Bulan", "Hari", "Jam", "Kategori", "Grup",
              "Membalas / akun asli", "Like", "Retweet", "Reply", "Quote", "Views", "Bookmark", "Interaksi",
              "Teks", "Link"])
    for i, r in enumerate(rows, 2):
        m, met = r["member"], r["metrics"]
        target = r["reply_to"] or r["orig_author"]
        link = (f"https://x.com/{r['orig_author']}/status/{r['orig_id']}" if r["kind"] == "Repost"
                else f"https://x.com/{m['handle_x']}/status/{r['id']}")
        d.append([
            d.c(m["handle_x"], text=True), d.c(m["panggilan"], text=True), d.c(m["tim"], text=True),
            d.c(r["wib"], fmt="yyyy-mm-dd hh:mm"), d.c(r["wib"].strftime("%Y-%m"), text=True),
            d.c(HARI[r["wib"].weekday()], text=True), d.c(r["wib"].hour), d.c(r["kind"], text=True),
            d.c(group_of(r["kind"]), text=True), d.c(handle_case.get(target.lower(), target), text=True),
            d.c(met["likes"], NUM), d.c(met["retweets"], NUM), d.c(met["replies"], NUM), d.c(met["quotes"], NUM),
            d.c(met["views"], NUM), d.c(met["bookmarks"], NUM), d.c(f"=K{i}+L{i}+M{i}+N{i}", NUM),
            d.c(r["text"], text=True),
            # URL sebagai teks: objek hyperlink per baris melampaui batas ~65 ribu hyperlink per sheet di Excel
            d.c(link, text=True, color=LINK),
        ])
    d.ws.auto_filter.ref = f"A1:S{last}"

    wb.calculation.fullCalcOnLoad = True
    wb.save(out)
    return out


def main():
    data_dir, out = Path(sys.argv[1]), Path(sys.argv[2])
    cutoff = datetime.strptime(sys.argv[3] if len(sys.argv) > 3 else "2025-10-22", "%Y-%m-%d")
    with open(BASE / "members.csv", newline="") as fh:
        members = [r for r in csv.DictReader(fh) if r["handle_x"]]
    if len(sys.argv) > 4:
        wanted = {h.lower() for h in sys.argv[4].split(",")}
        members = [m for m in members if m["handle_x"].lower() in wanted]
    print(build_workbook(members, data_dir, cutoff, out))


if __name__ == "__main__":
    main()
