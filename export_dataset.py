"""Dataset lengkap satu job dalam SATU file Excel: semua tweet member, mention, dan hashtag, satu baris per tweet.

export_jobs.py membuat Excel ringkasan di VPS (RAM 1 GB), sehingga baris mentahnya hanya disimpan sebagai CSV.gz.
Skrip ini menulis baris mentah itu ke Excel dan dijalankan di mesin dengan RAM cukup (data di-sync dari VPS).
Aturan privasinya sama dengan export_jobs.py:
  - fans individu (author_type 'lain'): tanpa handle, teks, link, atau follower pasti; id tweet di-hash dengan
    salt job (sama dengan kolom id_tweet di CSV.gz server); id penulis memakai hash dari xraw.
  - Tweet Member ditulis apa adanya (akun yang dibalas/di-repost tidak disamarkan): member tidak berinteraksi
    langsung dengan fans (keputusan user 2026-10-30).
  - retweet polos dibuang dari mention/hashtag. Di sheet Tweet Member, repost tetap ada tapi ditandai, karena
    metriknya milik tweet asli.
Satu tweet bisa menyebut beberapa member: di sheet Mention ia muncul sekali per member yang disebut.

Pakai: python3 export_dataset.py JOB_FILE DATA_ROOT OUT.xlsx
"""
import hashlib
import html
import json
import sys
from collections import Counter, defaultdict
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils import get_column_letter

from export_jobs import CENTER, HEAD_FILL, MAIN, anon_salt, f, label_of, phase_of, to_wib

MAX_ROWS = 1_048_575  # batas baris data per sheet Excel (di luar header)
MAX_CELL = 32_000     # batas teks per sel Excel 32.767
PCT, TIME, DATE = "0.00%", "yyyy-mm-dd hh:mm:ss", "yyyy-mm-dd"
AUTHOR_LABEL = {"member": "member", "fanbase": "fanbase", "official": "official", "lain": "fans/akun lain"}

# (kolom, lebar, format, arti) — dipakai untuk header sheet dan Kamus di sheet "Baca Dulu"
COMMON_TAIL = [
    ("like", 9, None, "Jumlah like saat diambil"),
    ("retweet", 9, None, "Jumlah retweet saat diambil"),
    ("reply", 9, None, "Jumlah reply saat diambil"),
    ("quote", 9, None, "Jumlah quote saat diambil"),
    ("views", 11, None, "Jumlah views saat diambil"),
    ("bookmark", 9, None, "Jumlah bookmark saat diambil"),
    ("engagement", 11, None, "like + retweet + reply + quote"),
    ("ER", 8, PCT, "engagement ÷ views (kosong bila views 0; di Tweet Member juga kosong untuk Repost)"),
]
COLS_MEMBER = [
    ("member", 11, None, "Panggilan member (lihat sheet Member)"),
    ("rank_2024", 8, None, "Peringkat SSK 2024 (kosong = tidak berperingkat)"),
    ("handle", 18, None, "Akun X member saat data diambil"),
    ("id_tweet", 20, None, "Id tweet (teks, supaya 19 digitnya tidak dibulatkan Excel)"),
    ("link", 22, None, "Tautan tweet"),
    ("waktu_wib", 19, TIME, "Waktu posting (WIB)"),
    ("tanggal_wib", 11, DATE, "Tanggal posting (WIB), untuk pivot harian"),
    ("fase", 12, None, "Pra-voting / Voting / Pasca-voting"),
    ("jenis", 20, None, "Original / Quote / Reply / Reply (thread sendiri) / Repost"),
    ("reply_ke", 16, None, "Akun yang dibalas (khusus Reply)"),
    ("akun_asli", 16, None, "Penulis tweet asli (khusus Repost)"),
    ("id_tweet_asli", 20, None, "Id tweet asli (khusus Repost)"),
    ("teks", 60, None, "Isi tweet (baris baru diganti spasi); Repost: teks tweet asli"),
    ("media", 12, None, "Jenis media terlampir"),
] + COMMON_TAIL + [
    ("dihitung_engagement", 11, None, "Tidak = Repost (metriknya milik tweet asli, jangan dijumlahkan)"),
    ("diambil_utc", 19, TIME, "Waktu data diambil (UTC)"),
]
COLS_SOURCE = [
    ("id_tweet", 20, None, "Id tweet; fans/akun lain: hash 't_…' (bukan id asli, cocok dengan CSV.gz server)"),
    ("link", 22, None, "Tautan tweet (kosong untuk fans/akun lain)"),
    ("waktu_wib", 19, TIME, "Waktu posting (WIB)"),
    ("tanggal_wib", 11, DATE, "Tanggal posting (WIB), untuk pivot harian"),
    ("fase", 12, None, "Pra-voting / Voting / Pasca-voting"),
    ("jenis", 9, None, "Original / Quote / Reply (Mention tanpa reply; Hashtag ikut reply; retweet polos dibuang)"),
    ("jenis_penulis", 13, None, "member / fanbase (dari dataset) / official / fans/akun lain"),
    ("penulis", 16, None, "Handle penulis (kosong untuk fans/akun lain)"),
    ("id_penulis", 19, None, "Handle huruf kecil, atau hash 'u_…' untuk fans/akun lain (tetap sama lintas sheet, "
     "untuk menghitung akun unik)"),
    ("follower_penulis", 11, None, "Follower penulis saat diambil (kosong untuk fans/akun lain)"),
    ("skala_follower", 11, None, "Kelompok follower penulis: <1rb, 1rb–10rb, 10rb–100rb, 100rb–1jt, ≥1jt"),
    ("penulis_terverifikasi", 10, None, "Centang biru (Ya/Tidak)"),
    ("tahun_akun_dibuat", 9, None, "Tahun akun penulis dibuat"),
    ("quote_dari", 16, None, "Akun yang di-quote (hanya akun publik: member/fanbase/official)"),
    ("menyebut", 24, None, "Akun publik yang disebut di tweet (handle fans dibuang)"),
    ("hashtag_tweet", 24, None, "Semua hashtag di tweet (huruf kecil)"),
    ("bahasa", 7, None, "Kode bahasa tweet menurut X"),
    ("media", 12, None, "Jenis media terlampir"),
] + COMMON_TAIL + [
    ("teks", 60, None, "Isi tweet (kosong untuk fans/akun lain)"),
    ("diambil_utc", 19, TIME, "Waktu data diambil (UTC)"),
]
REPLY_AT = 6  # posisi kolom reply_ke di sheet Hashtag, tepat setelah "jenis"
REPLY_COL = ("reply_ke", 16, None, "Akun publik yang dibalas (khusus Reply; kosong bila yang dibalas fans/akun lain)")
COLS_MENTION = [
    ("member_disebut", 11, None, "Member yang disebut (@handle sekarang/lama)"),
    ("rank_2024", 8, None, "Peringkat SSK 2024 member itu"),
] + COLS_SOURCE + [
    ("jumlah_member_disebut", 10, None, "Berapa member di dataset ini yang disebut tweet yang sama "
     "(tweet yang sama muncul sekali per member)"),
]
COLS_HASHTAG = [
    ("member", 11, None, "Member pemilik hashtag kampanye"),
    ("rank_2024", 8, None, "Peringkat SSK 2024 member itu"),
    ("hashtag_kampanye", 18, None, "Hashtag kampanye yang dicari (query lengkap di sheet Task)"),
] + COLS_SOURCE[:REPLY_AT] + [REPLY_COL] + COLS_SOURCE[REPLY_AT:] + [
    ("unik_per_member", 9, None, "1 = kemunculan pertama tweet ini untuk member itu; 0 = tweet yang sama sudah "
     "tercatat di hashtag lain member yang sama (filter 1 untuk menghitung tanpa dobel)"),
]


def clean(s):
    s = html.unescape(s or "")  # teks X memakai entitas HTML (&amp; &lt; &gt;)
    return ILLEGAL_CHARACTERS_RE.sub("", s.replace("\r", " ").replace("\n", " "))[:MAX_CELL]


def utc_naive(iso):
    return datetime.fromisoformat(iso).astimezone(timezone.utc).replace(tzinfo=None) if iso else None


def follower_scale(n):
    for limit, label in ((1_000, "<1rb"), (10_000, "1rb–10rb"), (100_000, "10rb–100rb"), (1_000_000, "100rb–1jt")):
        if n < limit:
            return label
    return "≥1jt"


def metrics(m, count=True):
    eng = m["likes"] + m["retweets"] + m["replies"] + m["quotes"]
    er = eng / m["views"] if count and m["views"] else None
    return [m["likes"], m["retweets"], m["replies"], m["quotes"], m["views"], m["bookmarks"], eng, er]


class Table:
    """Sheet data besar: header bergaya, baris berisi nilai polos (cepat), panel & filter di baris header."""

    def __init__(self, wb, title, cols):
        self.ws, self.cols, self.n = wb.create_sheet(title), cols, 0
        self.fmts = {i: c[2] for i, c in enumerate(cols) if c[2]}
        for i, c in enumerate(cols, 1):
            self.ws.column_dimensions[get_column_letter(i)].width = c[1]
        self.ws.freeze_panes = "A2"
        self.started = False

    def _cell(self, v, fmt=None, head=False):
        c = WriteOnlyCell(self.ws, value=v)
        if isinstance(v, str):
            c.data_type = "s"  # teks yang diawali '=' tetap teks
        if fmt:
            c.number_format = fmt
        if head:
            c.font, c.fill, c.alignment = f(True, "FFFFFF"), HEAD_FILL, CENTER
        return c

    def add(self, values):
        if not self.started:
            self.ws.append([self._cell(c[0], head=True) for c in self.cols])
            self.started = True
        if self.n >= MAX_ROWS:
            raise RuntimeError(f"sheet {self.ws.title} melewati batas baris Excel")
        row = []
        for i, v in enumerate(values):
            if v == "":
                v = None  # sel benar-benar kosong, bukan teks kosong
            fmt = self.fmts.get(i)
            if (fmt and v is not None) or (isinstance(v, str) and v.startswith("=")):
                v = self._cell(v, fmt)
            row.append(v)
        self.ws.append(row)
        self.n += 1

    def close(self):
        self.ws.auto_filter.ref = f"A1:{get_column_letter(len(self.cols))}{self.n + 1}"


class Notes:
    """Sheet teks/tabel kecil (Baca Dulu, Ringkasan, Member, Task)."""

    def __init__(self, wb, title, widths, freeze=None):
        self.ws = wb.create_sheet(title)
        for i, w in enumerate(widths, 1):
            self.ws.column_dimensions[get_column_letter(i)].width = w
        if freeze:
            self.ws.freeze_panes = freeze

    def text(self, s="", bold=False, italic=False):
        self.ws.append([self._c(s, bold, italic)] if s else [None])

    def _c(self, v, bold=False, italic=False):
        c = WriteOnlyCell(self.ws, value=v)
        if isinstance(v, str):
            c.data_type = "s"
        c.font = f(bold, italic=italic)
        return c

    def head(self, labels):
        cells = []
        for h in labels:
            c = WriteOnlyCell(self.ws, value=h)
            c.font, c.fill, c.alignment = f(True, "FFFFFF"), HEAD_FILL, CENTER
            cells.append(c)
        self.ws.append(cells)

    def row(self, values):
        self.ws.append([None if v == "" else self._c(v) if isinstance(v, str) else v for v in values])


def read_task(path, start, end):
    """Tweet unik satu file JSONL dalam [start, end) WIB tanpa retweet polos; versi terakhir per id menang."""
    rows = {}
    if path.exists():
        with open(path, encoding="utf-8") as fh:
            for line in fh:
                try:
                    r = json.loads(line)
                except json.JSONDecodeError:
                    continue  # baris terakhir bisa terpotong kalau proses mati saat menulis
                if r.get("kind") == "Repost":
                    continue
                wib = to_wib(r["created_at"])
                if start <= wib < end:
                    r["_wib"] = wib
                    rows[r["id"]] = r
    return sorted(rows.values(), key=lambda r: (r["_wib"], r["id"]))


def build_dataset(job, root, out):
    root, out = Path(root), Path(out)
    start = datetime.strptime(job["analysis_start_wib"], "%Y-%m-%d")
    end = datetime.strptime(job["analysis_end_wib"], "%Y-%m-%d") + timedelta(days=1)
    phases = job.get("phases", [])
    members = {m["key"]: m for m in job["members"]}
    order = sorted(members, key=lambda k: (members[k]["rank_2024"] or 999, members[k]["panggilan"]))
    tasks = {t["id"]: t for t in job["tasks"] if t["enabled"]}
    salt = anon_salt(root)
    if not salt:
        raise SystemExit("salt hash tidak ditemukan (DATA_ROOT/.hash_salt) — id tweet fans tidak boleh ditulis asli")

    def anon_id(tid):
        return "t_" + hashlib.sha256(f"{salt}:t{tid}".encode()).hexdigest()[:16]

    def jsonl(t):
        return root / t["mode"] / f"{label_of(t)}.jsonl"

    def source_cols(r, reply=False):
        fan = r["author_type"] == "lain"
        wib = r["_wib"]
        day = wib.strftime("%Y-%m-%d")
        followers = int(r.get("author_followers") or 0)
        year = (r.get("author_created") or "")[:4]
        cols = [anon_id(r["id"]) if fan else r["id"],
                "" if fan else f"https://x.com/i/status/{r['id']}",
                wib, wib.date(), phase_of(day, phases), r["kind"],
                AUTHOR_LABEL[r["author_type"]], "" if fan else r["author"], r["author_key"],
                None if fan else followers, follower_scale(followers),
                "Ya" if r.get("author_verified") else "Tidak", int(year) if year.isdigit() else None,
                r.get("quoted_author") or "", " ".join(r.get("mentions") or []),
                " ".join(r.get("hashtags") or []), r.get("lang") or "", ",".join(r.get("media") or [])] \
            + metrics(r["metrics"]) + ["" if fan else clean(r.get("text")), utc_naive(r.get("fetched_at"))]
        if reply:
            cols.insert(REPLY_AT, r.get("reply_to") or "")
        return cols

    wb = Workbook(write_only=True)
    readme = Notes(wb, "Baca Dulu", [26, 24, 90])
    summary = Notes(wb, "Ringkasan", [8, 12, 10, 10] + [13] * 14, freeze="C2")
    roster = Notes(wb, "Member", [12, 26, 12, 10, 8, 10, 18, 18, 18, 40, 12], freeze="B2")
    task_sheet = Notes(wb, "Task", [30, 9, 11, 22, 70, 9, 9, 10, 10, 10, 11], freeze="B2")
    t_member = Table(wb, "Tweet Member", COLS_MEMBER)
    t_mention = Table(wb, "Mention", COLS_MENTION)
    t_hashtag = Table(wb, "Hashtag", COLS_HASHTAG)
    stats = defaultdict(Counter)
    authors = defaultdict(lambda: defaultdict(set))
    task_rows = {}

    # 1) Tweet member
    profiles, member_windows = {}, {}
    for k in order:
        m = members[k]
        t = tasks.get(f"member:{k}")
        path = root / "member" / f"{m['handle']}_v2.json"
        if not t or not path.exists():
            continue
        raw = json.loads(path.read_text(encoding="utf-8"))
        profiles[k] = raw.get("profile") or {}
        member_windows[t["id"]] = len(raw.get("done") or [])
        rows = sorted((x for x in raw["tweets"].values() if start <= to_wib(x["created_at"]) < end),
                      key=lambda x: (x["created_at"], x["id"]))
        for x in rows:
            wib = to_wib(x["created_at"])
            repost = x["kind"] == "Repost"
            t_member.add([m["panggilan"], m["rank_2024"], m["handle"], x["id"], f"https://x.com/i/status/{x['id']}",
                          wib, wib.date(), phase_of(wib.strftime("%Y-%m-%d"), phases), x["kind"],
                          x.get("reply_to") or "", x.get("orig_author") or "", x.get("orig_id") or "",
                          clean(x.get("text")), ",".join(x.get("media") or [])]
                         + metrics(x["metrics"], count=not repost)
                         + ["Tidak" if repost else "Ya", utc_naive(x.get("fetched_at"))])
            stats[k]["member"] += 1
            stats[k]["member_utama"] += x["kind"] in MAIN
            stats[k]["member_reply"] += x["kind"].startswith("Reply")
            stats[k]["member_repost"] += repost
        task_rows[t["id"]] = len(rows)
    t_member.close()

    # 2) Mention: hitung dulu berapa member yang disebut tiap tweet, lalu tulis per member
    mention_tasks = [tasks[f"mention:{k}"] for k in order
                     if f"mention:{k}" in tasks and jsonl(tasks[f"mention:{k}"]).exists()]
    shared = Counter()
    for t in mention_tasks:
        shared.update({r["id"] for r in read_task(jsonl(t), start, end)})
    for t in mention_tasks:
        k = t["member"]
        rows = read_task(jsonl(t), start, end)
        for r in rows:
            t_mention.add([members[k]["panggilan"], members[k]["rank_2024"]] + source_cols(r) + [shared[r["id"]]])
            stats[k]["mention"] += 1
            stats[k]["mention_fans"] += r["author_type"] == "lain"
            stats[k]["mention_fanbase"] += r["author_type"] == "fanbase"
            stats[k]["mention_views"] += r["metrics"]["views"]
            authors[k]["mention"].add(r["author_key"])
        task_rows[t["id"]] = len(rows)
        del rows
    t_mention.close()

    # 3) Hashtag kampanye: satu baris per (hashtag, tweet); unik_per_member menandai kemunculan pertama
    seen = defaultdict(set)
    for k in order:
        for t in (x for x in job["tasks"] if x["enabled"] and x["mode"] == "hashtag" and x["member"] == k):
            if not jsonl(t).exists():
                continue
            rows = read_task(jsonl(t), start, end)
            for r in rows:
                first = r["id"] not in seen[k]
                seen[k].add(r["id"])
                t_hashtag.add([members[k]["panggilan"], members[k]["rank_2024"], "#" + t["tag"]] + source_cols(r, reply=True)
                              + [1 if first else 0])
                stats[k]["hashtag_baris"] += 1
                stats[k]["hashtag"] += first
                stats[k]["hashtag_views"] += r["metrics"]["views"] if first else 0
                authors[k]["hashtag"].add(r["author_key"])
            task_rows[t["id"]] = len(rows)
    t_hashtag.close()

    # 4) Ringkasan per member (angka dari baris dataset ini)
    sum_cols = ["Rank 2024", "Member", "Status", "Tim", "Tweet member (semua jenis)", "Post utama (Original+Quote)",
                "Reply member", "Repost member", "Mention", "Mention fans/akun lain", "Mention fanbase",
                "Mention akun unik", "Mention views", "Hashtag tweet unik", "Hashtag baris", "Hashtag akun unik",
                "Hashtag views (unik)"]
    summary.head(sum_cols)
    for k in order:
        m, s = members[k], stats[k]
        summary.row([m["rank_2024"], m["panggilan"], m["status"], m["tim"], s["member"], s["member_utama"],
                     s["member_reply"], s["member_repost"], s["mention"], s["mention_fans"], s["mention_fanbase"],
                     len(authors[k]["mention"]), s["mention_views"], s["hashtag"], s["hashtag_baris"],
                     len(authors[k]["hashtag"]), s["hashtag_views"]])
    tot = Counter()
    for s in stats.values():
        tot.update(s)
    summary.row(["", "TOTAL", "", "", tot["member"], tot["member_utama"], tot["member_reply"], tot["member_repost"],
                 tot["mention"], tot["mention_fans"], tot["mention_fanbase"], "", tot["mention_views"],
                 tot["hashtag"], tot["hashtag_baris"], "", tot["hashtag_views"]])

    # 5) Member & Task (konteks dan reproduksi)
    roster.head(["Member", "Nama", "Status", "Tim", "Rank 2024", "Key", "Handle X", "Handle lama", "Fanbase",
                 "Tagline/hashtag (dataset)", "Follower saat diambil"])
    for k in order:
        m = members[k]
        roster.row([m["panggilan"], m["nama"], m["status"], m["tim"], m["rank_2024"], k, m["handle"],
                    m["former_handle"], m["fanbase"], m["tagline_raw"],
                    profiles.get(k, {}).get("followers_count") or "–"])
    task_sheet.head(["Task", "Mode", "Member", "Hashtag", "Query", "Jendela (jam/hari)", "Pecah (jam)",
                     "Ambang pecah", "Jendela selesai", "Jendela terpotong", "Baris di dataset"])
    for t in job["tasks"]:
        if not t["enabled"]:
            continue
        st_path = root / t["mode"] / f"{label_of(t)}.state.json"
        wins = list(json.loads(st_path.read_text())["windows"].values()) if st_path.exists() else []
        if t["mode"] == "member":  # jendela member tercatat di "done" file _v2.json (tanpa tanda terpotong)
            done_n, cut = member_windows.get(t["id"]), None
        else:
            done_n, cut = (len(wins), sum(w["truncated"] for w in wins)) if wins else (None, None)
        window = f"{t['window_days']} hari" if "window_days" in t else f"{t.get('window_hours', '')} jam"
        task_sheet.row([t["id"], t["mode"], members[t["member"]]["panggilan"], ("#" + t["tag"]) if t.get("tag") else "",
                        t.get("query") or f"from:{t.get('handle', '')}", window, t.get("split_hours") or "",
                        t.get("split_threshold") or "", done_n, cut, task_rows.get(t["id"])])

    # 6) Baca Dulu (ditulis terakhir karena memuat jumlah baris; tetap sheet pertama)
    readme.text(f"Dataset X — {job.get('description', job['name'])}", bold=True)
    readme.text(f"Periode {job['analysis_start_wib']} s.d. {job['analysis_end_wib']} WIB · {len(order)} member · "
                f"dibuat {datetime.now():%Y-%m-%d %H:%M} dari data mentah job '{job['name']}'", italic=True)
    readme.text()
    readme.head(["Sheet", "Baris", "Isi"])
    for name, n, desc in [
        ("Ringkasan", len(order) + 1, "Jumlah per member (+1 baris TOTAL) dihitung dari baris di file ini"),
        ("Member", len(order), "Daftar member: handle sekarang & lama, fanbase, tagline/hashtag dari dataset SSK"),
        ("Task", len(task_rows), "Query X per task, ukuran jendela pencarian, dan kelengkapan (jendela terpotong = 0)"),
        ("Tweet Member", t_member.n, "Semua tweet akun member: Original, Quote, Reply, Reply (thread sendiri), Repost"),
        ("Mention", t_mention.n, "Tweet orang lain yang menyebut @member, tanpa reply & retweet polos; quote ikut. "
                                 "Satu baris per (member disebut, tweet)"),
        ("Hashtag", t_hashtag.n, "Tweet berisi hashtag kampanye + konteks JKT48/member/fanbase. "
                                 "Satu baris per (hashtag, tweet)"),
    ]:
        readme.row([name, n, desc])
    readme.text()
    for line in [
        "Catatan",
        "• Semua waktu dalam WIB (UTC+7). Like/retweet/views adalah angka saat diambil (Okt 2026), bukan saat periode.",
        "• Fase: " + "; ".join(f"{p['name']} {p['start']} s.d. {p['end']}" for p in phases) + ".",
        "• Fans individu & akun lain di luar daftar member/fanbase/official tampil anonim di sheet Mention & Hashtag: "
        "tanpa handle, teks, link, dan follower pasti. id_tweet-nya hash 't_…', id_penulis hash 'u_…' (tetap sama "
        "lintas sheet). Sheet Tweet Member ditulis apa adanya (member tidak berinteraksi langsung dengan fans).",
        "• Anonim di sini = pseudonim: waktu & metrik persis tetap ada untuk analisis, sehingga tweet aslinya masih "
        "bisa dicari ulang di X. Jangan bagikan baris fans satu per satu ke publik.",
        "• Tweet yang menyebut beberapa member muncul sekali per member di sheet Mention (lihat "
        "jumlah_member_disebut). Di sheet Hashtag, filter unik_per_member = 1 untuk menghitung tanpa dobel.",
        "• Repost di sheet Tweet Member membawa metrik tweet asli (dihitung_engagement = Tidak).",
        "• Search X tidak menjamin 100% lengkap: akun protected, tersaring, atau sudah dihapus tidak ikut.",
        "• Kalau Excel lambat: pakai filter kolom atau PivotTable; jangan sort seluruh sheet Mention sekaligus.",
    ]:
        readme.text(line, bold=line == "Catatan", italic=line.startswith("•"))
    readme.text()
    readme.text("Kamus kolom", bold=True)
    readme.head(["Sheet", "Kolom", "Arti"])
    for sheet, cols in [("Tweet Member", COLS_MEMBER), ("Mention", COLS_MENTION), ("Hashtag", COLS_HASHTAG)]:
        for c in cols:
            readme.row([sheet, c[0], c[3]])
    if job.get("events"):
        readme.text()
        readme.text("Tanggal penting", bold=True)
        for e in job["events"]:
            readme.text(f"{e['date']} — {e['name']}")

    out.parent.mkdir(parents=True, exist_ok=True)
    tmp = out.with_name(out.stem + ".tmp.xlsx")
    wb.save(tmp)
    tmp.replace(out)
    return {"file": str(out), "tweet_member": t_member.n, "mention": t_mention.n, "hashtag": t_hashtag.n,
            "per_member": {members[k]["panggilan"]: dict(stats[k]) for k in order}}


def main():
    job = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
    res = build_dataset(job, Path(sys.argv[2]), Path(sys.argv[3]))
    print(json.dumps({k: v for k, v in res.items() if k != "per_member"}, ensure_ascii=False))


if __name__ == "__main__":
    main()
