"""Ekspor data/<HANDLE>_v2.json ke Excel: Ringkasan, Bulanan, Reply, Pola Posting, Top 20, Data.

Pakai: python3 export_excel.py HANDLE [CUTOFF_YYYY-MM-DD]
"""
import json
import sys
from collections import Counter
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.formula import ArrayFormula

BASE = Path(__file__).parent
WIB = timedelta(hours=7)
HARI = ["Senin", "Selasa", "Rabu", "Kamis", "Jumat", "Sabtu", "Minggu"]
KINDS = ["Original", "Quote", "Reply", "Reply (thread sendiri)", "Repost"]
UTAMA, REPLY, REPOST = "Post utama", "Reply", "Repost"

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
SUB_FILL = PatternFill("solid", fgColor="D9E1F2")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(bottom=THIN)
NUM, PCT, DEC = "#,##0", "0.00%", "0.0"
GREY = "595959"


def f(bold=False, color="000000", size=10, italic=False):
    return Font(name=FONT, bold=bold, color=color, size=size, italic=italic)


def header(ws, row, labels, col=1):
    for i, h in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=h)
        c.font, c.fill = f(True, "FFFFFF"), HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = f(True, size=12)
    if sub:
        ws["A2"] = sub
        ws["A2"].font = f(italic=True, color=GREY)


def style_block(ws, r1, r2, c1, c2, fmt=NUM, total_row=None, fmts=None):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(r, c)
            cell.font = f(r == total_row)
            cell.number_format = (fmts or {}).get(c, fmt)
            cell.border = BOX
            if r == total_row:
                cell.fill = SUB_FILL


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def load(handle, cutoff):
    raw = json.loads((BASE / "data" / f"{handle}_v2.json").read_text())
    rows = []
    for t in raw["tweets"].values():
        utc = datetime.fromisoformat(t["created_at"]).replace(tzinfo=None)
        if utc < cutoff:
            continue
        rows.append(dict(t, wib=utc + WIB))
    rows.sort(key=lambda r: r["wib"])
    return raw["profile"], datetime.fromisoformat(raw["fetched_at"]).replace(tzinfo=None) + WIB, rows


def main():
    handle = sys.argv[1].lstrip("@")
    cutoff = datetime.strptime(sys.argv[2] if len(sys.argv) > 2 else "2025-10-22", "%Y-%m-%d")
    profile, fetched, rows = load(handle, cutoff)

    wb = Workbook()
    s = wb.active
    s.title = "Ringkasan"
    ws_mon = wb.create_sheet("Bulanan")
    ws_rep = wb.create_sheet("Reply")
    ws_pat = wb.create_sheet("Pola Posting")
    ws_top = wb.create_sheet("Top 20")
    ws = wb.create_sheet("Data")

    # ── Data ──────────────────────────────────────────────────────────
    cols = ["Tanggal (WIB)", "Bulan", "Hari", "Jam", "Kategori", "Grup", "Membalas / akun asli",
            "Like", "Retweet", "Reply", "Quote", "Views", "Bookmark", "Interaksi", "ER", "Media", "Teks", "Link"]
    header(ws, 1, cols)
    for i, r in enumerate(rows, 2):
        ws.cell(i, 1, r["wib"]).number_format = "yyyy-mm-dd hh:mm"
        ws.cell(i, 2, f'=TEXT(A{i},"yyyy-mm")')
        ws.cell(i, 3, HARI[r["wib"].weekday()])
        ws.cell(i, 4, f"=HOUR(A{i})")
        ws.cell(i, 5, r["kind"])
        ws.cell(i, 6, f'=IF(OR(E{i}="Original",E{i}="Quote"),"{UTAMA}",IF(E{i}="Repost","{REPOST}","{REPLY}"))')
        ws.cell(i, 7, r["reply_to"] or r["orig_author"])
        m = r["metrics"]
        for j, k in enumerate(["likes", "retweets", "replies", "quotes", "views", "bookmarks"], 8):
            ws.cell(i, j, m[k]).number_format = NUM
        ws.cell(i, 14, f"=H{i}+I{i}+J{i}+K{i}").number_format = NUM
        ws.cell(i, 15, f"=IF(L{i}>0,N{i}/L{i},0)").number_format = PCT
        ws.cell(i, 16, ",".join(r["media"]) or "teks")
        ws.cell(i, 17, r["text"].replace("\n", " "))
        link = (f"https://x.com/{r['orig_author']}/status/{r['orig_id']}" if r["kind"] == "Repost"
                else f"https://x.com/{handle}/status/{r['id']}")
        c = ws.cell(i, 18, link)
        c.hyperlink = link
        for j in range(1, 18):
            ws.cell(i, j).font = f()
        c.font = f(color="0563C1")
    last = len(rows) + 1
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:R{last}"
    widths(ws, [16, 9, 9, 6, 20, 11, 18, 9, 9, 8, 7, 11, 9, 10, 8, 10, 60, 30])

    def rng(col):
        return f"Data!${col}$2:${col}${last}"

    grp = lambda g: f'{rng("F")},"{g}"'
    knd = lambda k: f'{rng("E")},"{k}"'
    months = sorted({r["wib"].strftime("%Y-%m") for r in rows})
    m1, mlast = 5, 4 + len(months)
    tot = mlast + 1

    # ── Bulanan ───────────────────────────────────────────────────────
    title(ws_mon, f"Metrik bulanan @{handle}",
          f"Bulan pertama & terakhir tidak penuh: data {cutoff:%d %b %Y} – {fetched:%d %b %Y}.")
    # Tabel 1: jumlah per kategori
    ws_mon.cell(3, 1, "Jumlah tweet per kategori").font = f(True)
    header(ws_mon, 4, ["Bulan"] + KINDS + ["Total"])
    for i, mth in enumerate(months, m1):
        ws_mon.cell(i, 1, mth)
        for j, k in enumerate(KINDS, 2):
            ws_mon.cell(i, j, f'=COUNTIFS({knd(k)},{rng("B")},$A{i})')
        ws_mon.cell(i, 7, f"=SUM(B{i}:F{i})")
    ws_mon.cell(tot, 1, "Total")
    for j in range(2, 8):
        col = get_column_letter(j)
        ws_mon.cell(tot, j, f"=SUM({col}{m1}:{col}{mlast})")
    style_block(ws_mon, m1, tot, 1, 7, total_row=tot)

    # Tabel 2: performa post utama (original + quote)
    p0 = tot + 3
    ws_mon.cell(p0 - 1, 1, "Performa post utama (Original + Quote)").font = f(True)
    header(ws_mon, p0, ["Bulan", "Jumlah", "Total views", "Total like", "Total RT", "Total reply diterima",
                        "Total bookmark", "Rata-rata views", "Rata-rata like", "ER"])
    p1, plast = p0 + 1, p0 + len(months)
    ptot = plast + 1
    for i, mth in zip(range(p1, plast + 1), months):
        ws_mon.cell(i, 1, mth)
        crit = f'{grp(UTAMA)},{rng("B")},$A{i}'
        ws_mon.cell(i, 2, f"=COUNTIFS({crit})")
        for col, src in zip("CDEFG", "LHIJM"):
            ws_mon[f"{col}{i}"] = f"=SUMIFS({rng(src)},{crit})"
        ws_mon[f"H{i}"] = f"=IF(B{i}>0,C{i}/B{i},0)"
        ws_mon[f"I{i}"] = f"=IF(B{i}>0,D{i}/B{i},0)"
        ws_mon[f"J{i}"] = f"=IF(C{i}>0,SUMIFS({rng('N')},{crit})/C{i},0)"
    ws_mon.cell(ptot, 1, "Total")
    ws_mon.cell(ptot, 2, f"=SUM(B{p1}:B{plast})")
    for col in "CDEFG":
        ws_mon[f"{col}{ptot}"] = f"=SUM({col}{p1}:{col}{plast})"
    ws_mon[f"H{ptot}"] = f"=IF(B{ptot}>0,C{ptot}/B{ptot},0)"
    ws_mon[f"I{ptot}"] = f"=IF(B{ptot}>0,D{ptot}/B{ptot},0)"
    ws_mon[f"J{ptot}"] = f"=IF(C{ptot}>0,SUMIFS({rng('N')},{grp(UTAMA)})/C{ptot},0)"
    style_block(ws_mon, p1, ptot, 1, 10, total_row=ptot, fmts={10: PCT})

    # Tabel 3: performa reply
    r0 = ptot + 3
    ws_mon.cell(r0 - 1, 1, "Performa reply (ke akun lain + thread sendiri)").font = f(True)
    header(ws_mon, r0, ["Bulan", "Jumlah reply", "Total views", "Total like", "Rata-rata views",
                        "Rata-rata like", "Akun berbeda yang dibalas*"])
    r1, rlast = r0 + 1, r0 + len(months)
    rtot = rlast + 1
    per_month_targets = {}
    for r in rows:
        if r["kind"] == "Reply":
            per_month_targets.setdefault(r["wib"].strftime("%Y-%m"), set()).add(r["reply_to"].lower())
    for i, mth in zip(range(r1, rlast + 1), months):
        ws_mon.cell(i, 1, mth)
        crit = f'{grp(REPLY)},{rng("B")},$A{i}'
        ws_mon.cell(i, 2, f"=COUNTIFS({crit})")
        ws_mon[f"C{i}"] = f"=SUMIFS({rng('L')},{crit})"
        ws_mon[f"D{i}"] = f"=SUMIFS({rng('H')},{crit})"
        ws_mon[f"E{i}"] = f"=IF(B{i}>0,C{i}/B{i},0)"
        ws_mon[f"F{i}"] = f"=IF(B{i}>0,D{i}/B{i},0)"
        ws_mon.cell(i, 7, len(per_month_targets.get(mth, ()))).font = f(color="0000FF")
    ws_mon.cell(rtot, 1, "Total")
    for col in "BCD":
        ws_mon[f"{col}{rtot}"] = f"=SUM({col}{r1}:{col}{rlast})"
    ws_mon[f"E{rtot}"] = f"=IF(B{rtot}>0,C{rtot}/B{rtot},0)"
    ws_mon[f"F{rtot}"] = f"=IF(B{rtot}>0,D{rtot}/B{rtot},0)"
    all_targets = {r["reply_to"].lower() for r in rows if r["kind"] == "Reply"}
    ws_mon.cell(rtot, 7, len(all_targets))
    style_block(ws_mon, r1, rtot, 1, 7, total_row=rtot)
    for i in range(r1, rtot + 1):
        ws_mon.cell(i, 7).font = f(i == rtot, color="0000FF")
    ws_mon.cell(rtot + 1, 1, "* Dihitung saat ekspor (bukan rumus); tidak termasuk reply ke thread sendiri.").font = \
        f(italic=True, color=GREY)
    widths(ws_mon, [10, 11, 13, 14, 13, 14, 14, 13, 12, 10])
    ws_mon.freeze_panes = "B5"

    ch = LineChart()
    ch.title, ch.y_axis.title, ch.height, ch.width = "Post utama: rata-rata views & like per post", "Views", 8, 20
    ch.add_data(Reference(ws_mon, min_col=8, min_row=p0, max_row=plast), titles_from_data=True)
    ch.set_categories(Reference(ws_mon, min_col=1, min_row=p1, max_row=plast))
    ch2 = LineChart()
    ch2.add_data(Reference(ws_mon, min_col=9, min_row=p0, max_row=plast), titles_from_data=True)
    ch2.y_axis.axId, ch2.y_axis.title, ch2.y_axis.crosses = 200, "Like", "max"
    ch += ch2
    ws_mon.add_chart(ch, "L4")
    bc = BarChart()
    bc.type, bc.grouping, bc.overlap = "col", "stacked", 100
    bc.title, bc.height, bc.width = "Komposisi tweet per bulan", 8, 20
    bc.add_data(Reference(ws_mon, min_col=2, max_col=6, min_row=4, max_row=mlast), titles_from_data=True)
    bc.set_categories(Reference(ws_mon, min_col=1, min_row=m1, max_row=mlast))
    ws_mon.add_chart(bc, "L22")

    # ── Reply: akun yang paling sering dibalas ────────────────────────
    title(ws_rep, f"Siapa yang paling sering dibalas @{handle}?", "Reply ke thread sendiri tidak dihitung.")
    header(ws_rep, 4, ["#", "Akun", "Jumlah reply", "% dari semua reply", "Rata-rata views reply", "Reply terakhir"])
    targets = Counter(r["reply_to"] for r in rows if r["kind"] == "Reply")
    top_targets = targets.most_common(30)
    n_last = 4 + len(top_targets)
    for n, (acct, _) in enumerate(top_targets, 1):
        i = n + 4
        crit = f'{knd("Reply")},{rng("G")},$B{i}'
        ws_rep.cell(i, 1, n)
        ws_rep.cell(i, 2, acct)
        ws_rep.cell(i, 3, f"=COUNTIFS({crit})")
        ws_rep.cell(i, 4, f'=IF(COUNTIFS({knd("Reply")})>0,C{i}/COUNTIFS({knd("Reply")}),0)')
        ws_rep.cell(i, 5, f"=IF(C{i}>0,SUMIFS({rng('L')},{crit})/C{i},0)")
        ws_rep.cell(i, 6, f'=_xlfn.MAXIFS({rng("A")},{crit})')
    style_block(ws_rep, 5, max(5, n_last), 1, 6, fmts={4: "0.0%", 6: "yyyy-mm-dd"})
    ws_rep.cell(n_last + 2, 1, f"Urutan 30 besar dibuat saat ekspor dari {len(targets)} akun berbeda.").font = \
        f(italic=True, color=GREY)
    widths(ws_rep, [4, 22, 12, 12, 14, 14])

    # ── Pola Posting (post utama) ─────────────────────────────────────
    title(ws_pat, "Kapan post utama paling ramai? (Original + Quote, jam WIB)")
    header(ws_pat, 3, ["Hari", "Jumlah post", "Rata-rata views", "Rata-rata like"])
    for i, h in enumerate(HARI, 4):
        ws_pat.cell(i, 1, h)
        crit = f'{grp(UTAMA)},{rng("C")},$A{i}'
        ws_pat.cell(i, 2, f"=COUNTIFS({crit})")
        ws_pat.cell(i, 3, f"=IF(B{i}>0,SUMIFS({rng('L')},{crit})/B{i},0)")
        ws_pat.cell(i, 4, f"=IF(B{i}>0,SUMIFS({rng('H')},{crit})/B{i},0)")
    header(ws_pat, 3, ["Jam", "Jumlah post", "Rata-rata views", "Rata-rata like"], col=6)
    for i, hr in enumerate(range(24), 4):
        ws_pat.cell(i, 6, hr)
        crit = f'{grp(UTAMA)},{rng("D")},$F{i}'
        ws_pat.cell(i, 7, f"=COUNTIFS({crit})")
        ws_pat.cell(i, 8, f"=IF(G{i}>0,SUMIFS({rng('L')},{crit})/G{i},0)")
        ws_pat.cell(i, 9, f"=IF(G{i}>0,SUMIFS({rng('H')},{crit})/G{i},0)")
    style_block(ws_pat, 4, 10, 1, 4)
    style_block(ws_pat, 4, 27, 6, 9)
    widths(ws_pat, [10, 11, 15, 14, 3, 6, 11, 15, 14])
    hc = BarChart()
    hc.title, hc.height, hc.width, hc.legend = "Jumlah post utama per jam (WIB)", 7, 16, None
    hc.add_data(Reference(ws_pat, min_col=7, min_row=3, max_row=27), titles_from_data=True)
    hc.set_categories(Reference(ws_pat, min_col=6, min_row=4, max_row=27))
    ws_pat.add_chart(hc, "K3")

    # ── Top 20 post utama ─────────────────────────────────────────────
    title(ws_top, "20 post utama dengan views tertinggi (Original + Quote)")
    header(ws_top, 3, ["#", "Tanggal (WIB)", "Kategori", "Views", "Like", "Retweet", "Reply", "Bookmark", "ER",
                       "Teks", "Link"])
    main_rows = [(i, r) for i, r in enumerate(rows, 2) if r["kind"] in ("Original", "Quote")]
    for n, (src, r) in enumerate(sorted(main_rows, key=lambda x: -x[1]["metrics"]["views"])[:20], 1):
        i = n + 3
        ws_top.cell(i, 1, n)
        for col, s_col in zip("BCDEFGHIJK", "AELHIJMOQR"):
            ws_top[f"{col}{i}"] = f"=Data!{s_col}{src}"
        style_block(ws_top, i, i, 1, 11, fmts={2: "yyyy-mm-dd hh:mm", 3: "@", 9: PCT, 10: "@", 11: "@"})
        for c in range(2, 12):
            ws_top.cell(i, c).font = f(color="008000")
    ws_top.cell(25, 1, "Urutan dibuat saat ekspor; angka hijau = tautan ke sheet Data.").font = f(italic=True, color=GREY)
    widths(ws_top, [4, 16, 10, 11, 9, 9, 8, 10, 8, 50, 30])

    # ── Ringkasan ─────────────────────────────────────────────────────
    s["A1"] = f"{profile['name']} (@{handle}) — metrik X"
    s["A1"].font = f(True, size=14)
    s["A2"] = f"Periode {cutoff:%d %b %Y} – {fetched:%d %b %Y} · diambil {fetched:%d %b %Y %H:%M} WIB via twitter-cli"
    s["A2"].font = f(italic=True, color=GREY)

    def section(row, text, ncol):
        s.cell(row, 1, text).font = f(True)
        for c in range(1, ncol + 1):
            s.cell(row, c).fill = SUB_FILL

    section(4, "Profil (saat diambil)", 2)
    joined = datetime.strptime(profile["created_at"], "%a %b %d %H:%M:%S %z %Y").strftime("%d %b %Y")
    prof = [("Follower", profile["followers_count"]), ("Following", profile["following_count"]),
            ("Total post sepanjang masa", profile["tweets_count"]), ("Bergabung", joined)]
    for i, (k, v) in enumerate(prof, 5):
        s.cell(i, 1, k).font = f()
        c = s.cell(i, 2, v)
        c.font, c.number_format, c.alignment = f(color="0000FF"), NUM, Alignment(horizontal="right")
    s["B5"].comment = Comment("Hanya kondisi saat diambil; X tidak menyediakan riwayat follower.", "Claude")

    section(10, "Per kategori (1 tahun)", 7)
    header(s, 11, ["Kategori", "Jumlah", "% dari total", "Total views", "Rata-rata views", "Rata-rata like", "ER"])
    for i, k in enumerate(KINDS, 12):
        s.cell(i, 1, k)
        s.cell(i, 2, f"=COUNTIFS({knd(k)})")
        s.cell(i, 3, f"=IF($B$17>0,B{i}/$B$17,0)")
        if k == "Repost":
            for c in range(4, 8):
                s.cell(i, c, "–").alignment = Alignment(horizontal="right")
            s.cell(i, 4).comment = Comment("Like/views repost milik tweet asli, jadi tidak dihitung.", "Claude")
            continue
        s.cell(i, 4, f"=SUMIFS({rng('L')},{knd(k)})")
        s.cell(i, 5, f"=IF(B{i}>0,D{i}/B{i},0)")
        s.cell(i, 6, f"=IF(B{i}>0,SUMIFS({rng('H')},{knd(k)})/B{i},0)")
        s.cell(i, 7, f"=IF(D{i}>0,SUMIFS({rng('N')},{knd(k)})/D{i},0)")
    s.cell(17, 1, "Total")
    s.cell(17, 2, "=SUM(B12:B16)")
    s.cell(17, 3, "=SUM(C12:C16)")
    s.cell(17, 4, "=SUM(D12:D15)")
    style_block(s, 12, 17, 1, 7, total_row=17, fmts={3: "0.0%", 7: PCT})

    section(19, "Post utama (Original + Quote)", 2)
    M = "'Bulanan'!"
    stats = [
        ("Jumlah post utama", f"={M}B{ptot}", NUM),
        ("Total views", f"={M}C{ptot}", NUM),
        ("Total like", f"={M}D{ptot}", NUM),
        ("Total retweet", f"={M}E{ptot}", NUM),
        ("Total reply diterima", f"={M}F{ptot}", NUM),
        ("Total bookmark", f"={M}G{ptot}", NUM),
        ("Rata-rata views per post", f"={M}H{ptot}", NUM),
        ("Median views per post", None, NUM),
        ("Rata-rata like per post", f"={M}I{ptot}", NUM),
        ("Engagement rate", f"={M}J{ptot}", PCT),
        ("Post utama per bulan (rata-rata)", f"=IF(COUNTA({M}A{p1}:A{plast})>0,B20/COUNTA({M}A{p1}:A{plast}),0)", DEC),
        ("Bulan dengan views rata-rata tertinggi",
         f"=INDEX({M}A{p1}:A{plast},MATCH(MAX({M}H{p1}:H{plast}),{M}H{p1}:H{plast},0))", "@"),
    ]
    for i, (k, formula, fmt) in enumerate(stats, 20):
        s.cell(i, 1, k).font = f()
        c = s.cell(i, 2, formula)
        if formula is None:
            # median bersyarat butuh array formula agar jalan di semua versi Excel
            s.cell(i, 2).value = ArrayFormula(f"B{i}", f'=MEDIAN(IF({rng("F")}="{UTAMA}",{rng("L")}))')
            c = s.cell(i, 2)
        c.font, c.number_format, c.alignment = f(color="008000"), fmt, Alignment(horizontal="right")

    notes = [
        "Catatan",
        "• Post utama = Original + Quote. Reply = balasan ke akun lain; 'thread sendiri' = membalas tweet sendiri.",
        "• Repost tidak dihitung di metrik karena like/views-nya milik tweet asli.",
        "• Like, views, dll. adalah angka saat diambil, bukan saat tweet diposting.",
        "• ER = (like + retweet + reply + quote) ÷ views.",
        "• Semua data dari search 'from:akun' per minggu; search X bisa sedikit tidak lengkap.",
        "• Biru = angka tetap; hijau = rumus yang menautkan sheet lain; selebihnya rumus dari sheet Data.",
    ]
    for i, n in enumerate(notes, 33):
        s.cell(i, 1, n).font = f(i == 33, color="000000" if i == 33 else GREY, italic=i != 33)
    widths(s, [38, 14, 12, 14, 15, 14, 9])

    wb.calculation.fullCalcOnLoad = True
    out = BASE / f"{handle}_metrik_X.xlsx"
    wb.save(out)
    print(out, len(rows), "baris")


if __name__ == "__main__":
    main()
