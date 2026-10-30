"""Excel & ringkasan untuk hasil job_worker (mode member, mention, hashtag/tagline).

Data mention/hashtag bisa ratusan ribu tweet per task, jadi dibaca per task secara streaming ke objek ringkas
(tanpa menahan JSON mentah) dan dibuang sebelum task berikutnya. Angka di Excel dihitung di Python saat ekspor
(nilai, bukan rumus) supaya file tetap kecil. Retweet polos dibuang (metriknya milik tweet asli).
Fans individu (author_type 'lain') tidak pernah ditampilkan dengan handle, teks, link, atau id tweet asli.
Data mentah ringkas per ekspor: <root>/exports/<mode>_<nama-file-excel>.csv.gz (tidak dikirim ke Telegram).

Pakai manual: python3 export_jobs.py JOB_FILE DATA_ROOT OUT.xlsx
"""
import csv
import gzip
import hashlib
import heapq
import json
import os
import statistics
import sys
from collections import Counter, defaultdict
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

WIB = timezone(timedelta(hours=7))
MAIN = ("Original", "Quote")
FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
NUM, PCT, DEC = "#,##0", "0.0%", "0.0"
_FONTS = {}
TOP_N = 100
AUTHOR_TYPES = ["fanbase", "lain", "member", "official"]
ANON_TEXT = "(teks fans individu tidak disimpan)"


def f(bold=False, color="000000", italic=False, size=10):
    key = (bold, color, italic, size)
    if key not in _FONTS:
        _FONTS[key] = Font(name=FONT, bold=bold, color=color, italic=italic, size=size)
    return _FONTS[key]


def to_wib(iso):
    return datetime.fromisoformat(iso).astimezone(WIB).replace(tzinfo=None)


def phase_of(day, phases):
    for p in phases:
        if p["start"] <= day <= p["end"]:
            return p["name"]
    return ""


def days_in(p):
    return (datetime.fromisoformat(p["end"]) - datetime.fromisoformat(p["start"])).days + 1


def label_of(t):
    return t["id"].replace(":", "__")


def anon_salt(root):
    path = Path(root) / ".hash_salt"
    return os.environ.get("HASH_SALT", "").strip() or (path.read_text().strip() if path.exists() else "")


# ── Membaca hasil per task (streaming, ringkas) ────────────────────────
class Row:
    __slots__ = ("id", "wib", "kind", "author", "author_key", "author_type", "followers", "likes", "retweets",
                 "replies", "quotes", "views", "bookmarks", "lang", "media", "hashtags", "text")

    def __init__(self, r):
        m = r["metrics"]
        self.id = r["id"]
        self.wib = to_wib(r["created_at"])
        self.kind = sys.intern(r["kind"])
        self.author_type = sys.intern(r["author_type"])
        public = self.author_type != "lain"
        self.author = sys.intern(r.get("author") or "") if public else ""
        self.author_key = sys.intern(r["author_key"])
        self.followers = int(r.get("author_followers") or 0)
        self.likes, self.retweets, self.replies = m["likes"], m["retweets"], m["replies"]
        self.quotes, self.views, self.bookmarks = m["quotes"], m["views"], m["bookmarks"]
        self.lang = sys.intern(r.get("lang") or "")
        self.media = sys.intern(",".join(r.get("media") or []))
        self.hashtags = " ".join(r.get("hashtags") or [])
        self.text = ILLEGAL_CHARACTERS_RE.sub("", (r.get("text") or "").replace("\n", " "))[:280] if public else ""


def task_ran(root, t):
    return (Path(root) / t["mode"] / f"{label_of(t)}.state.json").exists()


def load_task(root, t, start, end):
    """Tweet unik task dalam [start, end) WIB, tanpa retweet polos; versi terakhir per id menang."""
    root = Path(root)
    path = root / t["mode"] / f"{label_of(t)}.jsonl"
    rows = {}
    if path.exists():
        with open(path, encoding="utf-8") as fh:
            for line in fh:
                try:
                    r = json.loads(line)
                except json.JSONDecodeError:
                    continue  # baris terakhir bisa terpotong kalau proses mati saat menulis
                if r.get("kind") == "Repost":
                    continue
                row = Row(r)
                if start <= row.wib < end:
                    rows[row.id] = row
    state_path = root / t["mode"] / f"{label_of(t)}.state.json"
    st = json.loads(state_path.read_text()) if state_path.exists() else {"windows": {}}
    return list(rows.values()), st


def summarize(rows, phases):
    n = len(rows)
    by_type = Counter(r.author_type for r in rows)
    out = {"tweet": n, "akun unik": len({r.author_key for r in rows}),
           "total like": sum(r.likes for r in rows), "total views": sum(r.views for r in rows),
           "total RT": sum(r.retweets for r in rows),
           "median like": statistics.median([r.likes for r in rows]) if rows else 0,
           "median views": statistics.median([r.views for r in rows]) if rows else 0}
    for a in AUTHOR_TYPES:
        out[f"tweet {a}"] = by_type.get(a, 0)
    out["porsi fanbase"] = by_type.get("fanbase", 0) / n if n else 0
    for p in phases:
        pr = [r for r in rows if phase_of(r.wib.strftime("%Y-%m-%d"), [p])]
        out[f"{p['name']}: tweet/hari"] = len(pr) / days_in(p)
        out[f"{p['name']}: akun unik"] = len({r.author_key for r in pr})
    return out


def member_summary(path, start, end, phases):
    if not path.exists():
        return None
    raw = json.loads(path.read_text())
    rows = [t for t in raw["tweets"].values() if start <= to_wib(t["created_at"]) < end]
    main = [t for t in rows if t["kind"] in MAIN]
    out = {"follower (saat diambil)": (raw.get("profile") or {}).get("followers_count", 0),
           "post utama": len(main), "reply": sum(t["kind"].startswith("Reply") for t in rows),
           "repost": sum(t["kind"] == "Repost" for t in rows),
           "median like post": statistics.median([t["metrics"]["likes"] for t in main]) if main else 0,
           "median views post": statistics.median([t["metrics"]["views"] for t in main]) if main else 0}
    for p in phases:
        pm = [t for t in main if phase_of(to_wib(t["created_at"]).strftime("%Y-%m-%d"), [p])]
        out[f"{p['name']}: post/minggu"] = 7 * len(pm) / days_in(p)
        out[f"{p['name']}: median like"] = statistics.median([t["metrics"]["likes"] for t in pm]) if pm else 0
    return out


# ── Penulisan workbook ────────────────────────────────────────────────
class Sheet:
    def __init__(self, wb, title, widths):
        self.ws = wb.create_sheet(title)
        for i, wdt in enumerate(widths, 1):
            self.ws.column_dimensions[get_column_letter(i)].width = wdt

    def cell(self, v, fmt=None, bold=False, color="000000", fill=None, align=None, italic=False):
        c = WriteOnlyCell(self.ws, value=v)
        if isinstance(v, str):
            c.data_type = "s"  # teks yang diawali '=' tetap teks
        c.font = f(bold, color, italic)
        if fmt:
            c.number_format = fmt
        if fill:
            c.fill = fill
        if align:
            c.alignment = align
        return c

    def row(self, values, fmts=None):
        fmts = fmts or {}
        self.ws.append([self.cell(v, fmts.get(i, NUM if isinstance(v, (int, float)) else None))
                        for i, v in enumerate(values)])

    def header(self, labels):
        self.ws.append([self.cell(h, bold=True, color="FFFFFF", fill=HEAD_FILL, align=CENTER) for h in labels])

    def text(self, s, bold=False, italic=False, color="000000"):
        self.ws.append([self.cell(s, bold=bold, italic=italic, color=color)])


def fmt_for(key):
    if key.startswith("porsi"):
        return PCT
    if "/hari" in key or "/minggu" in key:
        return DEC
    return NUM


def selected_tasks(job, root, task_ids):
    """Task yang diminta DAN sudah pernah dijalankan (punya state); task lain tidak ditulis sebagai angka 0."""
    wanted = set(task_ids) if task_ids else None
    out = []
    for t in job["tasks"]:
        if wanted is not None and t["id"] not in wanted:
            continue
        if t["mode"] == "member":
            if (Path(root) / "member" / f"{t['handle']}_v2.json").exists():
                out.append(t)
        elif task_ran(root, t):
            out.append(t)
    return out


def build_job_workbook(job, root, out, task_ids=None):
    root, out = Path(root), Path(out)
    start = datetime.strptime(job["analysis_start_wib"], "%Y-%m-%d")
    end = datetime.strptime(job["analysis_end_wib"], "%Y-%m-%d") + timedelta(days=1)
    phases = job.get("phases", [])
    members = {m["key"]: m for m in job["members"]}
    tasks = selected_tasks(job, root, task_ids)
    exports = root / "exports"
    exports.mkdir(parents=True, exist_ok=True)
    salt = anon_salt(root)
    anon_id = (lambda tid: "t_" + hashlib.sha256(f"{salt}:t{tid}".encode()).hexdigest()[:16]) if salt \
        else (lambda tid: "")

    mention_sum, hashtag_rows, daily = {}, [], defaultdict(Counter)
    tag_ids, tag_authors, tag_views = defaultdict(set), defaultdict(set), defaultdict(int)
    tops = {"mention": [], "hashtag": []}
    completeness, writers = [], {}
    try:
        for t in tasks:
            if t["mode"] == "member":
                continue
            rows, st = load_task(root, t, start, end)
            wins = list(st["windows"].values())
            completeness.append([t["id"], len(wins), sum(x["truncated"] for x in wins),
                                 sum(x["pages"] for x in wins), len(rows)])
            mode = "mention" if t["mode"] == "mention" else "hashtag"
            s = summarize(rows, phases)
            if mode == "mention":
                mention_sum[t["member"]] = s
            else:
                hashtag_rows.append((t, s))
            col = f"{'mention' if mode == 'mention' else '#' + t.get('tag', '')} {members[t['member']]['panggilan']}"
            for r in rows:
                daily[r.wib.strftime("%Y-%m-%d")][col] += 1
                if mode == "hashtag":
                    k = t["member"]
                    if r.id not in tag_ids[k]:  # tweet berisi 2 hashtag member yang sama dihitung sekali
                        tag_ids[k].add(r.id)
                        tag_views[k] += r.views
                    tag_authors[k].add(r.author_key)
                item = (r.views, r.id, t["member"], t.get("tag", ""), r)
                heap = tops[mode]
                if len(heap) < TOP_N:
                    heapq.heappush(heap, item)
                elif item[:2] > heap[0][:2]:
                    heapq.heapreplace(heap, item)
            if mode not in writers:
                tmp = exports / f"{mode}_{out.stem}.csv.gz.tmp"
                fh = gzip.open(tmp, "wt", encoding="utf-8", newline="")
                wr = csv.writer(fh)
                wr.writerow(["task", "member", "tag", "id_tweet", "waktu_wib", "kind", "author", "author_key",
                             "author_type", "author_followers", "likes", "retweets", "replies", "quotes", "views",
                             "bookmarks", "lang", "media", "hashtags", "text"])
                writers[mode] = (fh, wr, tmp)
            wr = writers[mode][1]
            for r in rows:
                fan = r.author_type == "lain"
                wr.writerow([t["id"], t["member"], t.get("tag", ""), anon_id(r.id) if fan else r.id,
                             r.wib.isoformat(), r.kind, r.author, r.author_key, r.author_type,
                             "" if fan else r.followers, r.likes, r.retweets, r.replies, r.quotes, r.views,
                             r.bookmarks, r.lang, r.media, r.hashtags, r.text])
            del rows  # bebaskan memori sebelum task berikutnya
    finally:
        for fh, _, tmp in writers.values():
            fh.close()
            os.replace(tmp, str(tmp)[:-len(".tmp")])

    member_sum = {t["member"]: member_summary(root / "member" / f"{t['handle']}_v2.json", start, end, phases)
                  for t in tasks if t["mode"] == "member"}

    wb = Workbook(write_only=True)
    keys = sorted(members, key=lambda k: (members[k]["rank_2024"] or 999, members[k]["panggilan"]))

    # Ringkasan per member
    s = Sheet(wb, "Ringkasan", [8, 13, 10, 10] + [14] * 30)
    s.ws.freeze_panes = "C3"
    s.text(f"{job.get('description', job['name'])} — periode {job['analysis_start_wib']} s.d. "
           f"{job['analysis_end_wib']} WIB", bold=True)
    m_keys = next((list(v) for v in member_sum.values() if v), [])
    x_keys = list(next(iter(mention_sum.values()), {}))
    has_tag = bool(hashtag_rows)
    s.header(["Rank SSK 2024", "Member", "Status", "Tim"] + [f"Member · {k}" for k in m_keys]
             + [f"Mention · {k}" for k in x_keys]
             + (["Hashtag · tweet", "Hashtag · akun unik", "Hashtag · views"] if has_tag else []))
    for k in keys:
        m = members[k]
        ms, xs = member_sum.get(k) or {}, mention_sum.get(k) or {}
        ran_tag = any(t["member"] == k for t, _ in hashtag_rows)
        vals = [m["rank_2024"] or "–", m["panggilan"], m["status"], m["tim"]]
        vals += [ms.get(c, "–") for c in m_keys] + [xs.get(c, "–") for c in x_keys]
        if has_tag:
            vals += ([len(tag_ids[k]), len(tag_authors[k]), tag_views[k]] if ran_tag else ["–", "–", "–"])
        fmts = {4 + i: fmt_for(c) for i, c in enumerate(m_keys)}
        fmts.update({4 + len(m_keys) + i: fmt_for(c) for i, c in enumerate(x_keys)})
        s.row(vals, fmts)

    # Hashtag
    hs = Sheet(wb, "Hashtag", [12, 22, 50] + [14] * 30)
    hkeys = list(hashtag_rows[0][1]) if hashtag_rows else []
    hs.header(["Member", "Hashtag", "Query"] + hkeys)
    for t, h in sorted(hashtag_rows, key=lambda x: -x[1]["tweet"]):
        hs.row([members[t["member"]]["panggilan"], "#" + t.get("tag", ""), t["query"]] + [h[c] for c in hkeys],
               {3 + i: fmt_for(c) for i, c in enumerate(hkeys)})

    # Harian
    cols = sorted({c for d in daily.values() for c in d})
    hr = Sheet(wb, "Harian", [12] + [12] * len(cols))
    hr.ws.freeze_panes = "B2"
    hr.header(["Tanggal (WIB)"] + cols)
    events = {e["date"]: e["name"] for e in job.get("events", [])}
    day = start
    while day < end:
        d = day.strftime("%Y-%m-%d")
        hr.row([d + (f" · {events[d]}" if d in events else "")] + [daily[d].get(c, 0) for c in cols])
        day += timedelta(days=1)

    # Top tweet: fans individu tanpa handle/teks/link/follower
    for mode, title in [("mention", "Top Mention"), ("hashtag", "Top Hashtag")]:
        tp = Sheet(wb, title, [5, 16, 12, 16, 12, 18, 10, 10, 12, 60, 40])
        tp.header(["#", "Waktu (WIB)", "Member", "Hashtag", "Jenis penulis", "Penulis", "Follower penulis",
                   "Like", "Views", "Teks", "Link"])
        for i, (_, tid, mk, tag, r) in enumerate(sorted(tops[mode], key=lambda x: (-x[0], x[1])), 1):
            fan = r.author_type == "lain"
            tp.row([i, r.wib.strftime("%Y-%m-%d %H:%M"), members[mk]["panggilan"], ("#" + tag) if tag else "",
                    r.author_type, "fans (anonim)" if fan else r.author, "" if fan else r.followers, r.likes,
                    r.views, ANON_TEXT if fan else r.text, "" if fan else f"https://x.com/i/status/{tid}"])

    # Catatan & kelengkapan
    c = Sheet(wb, "Catatan", [40, 12, 12, 12, 12])
    notes = [
        "Catatan",
        "• Angka di file ini dihitung saat ekspor (nilai, bukan rumus) karena volume data mention/hashtag besar.",
        "• Mention = tweet orang lain yang menyebut @member (handle sekarang + lama), tanpa reply dan tanpa "
        "retweet polos; quote dihitung.",
        "• Hashtag dicari bersama konteks (JKT48/nama/@member/@fanbase) supaya hashtag umum (#CHANCE dll.) tidak "
        "tercampur tweet di luar JKT48. Lihat kolom Query di sheet Hashtag. Ringkasan per member menghitung "
        "tweet & akun unik lintas semua hashtag member itu tanpa dobel.",
        "• Jenis penulis: member / fanbase (dari dataset) / official / lain = fans individu & akun lain. "
        "Fans individu anonim: tanpa handle, teks, link, atau id tweet asli.",
        "• Like/views adalah angka saat diambil (tweet 2023–2024 sudah matang). Search X tidak menjamin 100% "
        "lengkap; akun protected/tersaring/terhapus tidak ikut.",
        f"• Data mentah ringkas: exports/mention_{out.stem}.csv.gz dan exports/hashtag_{out.stem}.csv.gz di server.",
        "• Task yang belum dijalankan tidak ditampilkan (–), bukan 0.",
        "",
        "Kelengkapan per task",
    ]
    for n in notes:
        c.text(n, bold=n in ("Catatan", "Kelengkapan per task"), italic=n.startswith("•"))
    c.header(["Task", "Jendela selesai", "Jendela terpotong", "Halaman", "Tweet unik"])
    for row in completeness:
        c.row(row)
    if job.get("events"):
        c.text("")
        c.text("Tanggal penting", bold=True)
        for e in job["events"]:
            c.text(f"{e['date']} — {e['name']}")

    wb.save(out)
    return out


def _id(x):
    return f"{x:,.0f}".replace(",", ".")


def _count_task(root, t, start, end):
    """Hitung tweet & akun unik secara streaming (hanya id & author_key yang ditahan)."""
    ids, authors = {}, set()
    path = Path(root) / t["mode"] / f"{label_of(t)}.jsonl"
    if path.exists():
        with open(path, encoding="utf-8") as fh:
            for line in fh:
                try:
                    r = json.loads(line)
                except json.JSONDecodeError:
                    continue
                if r.get("kind") == "Repost" or not start <= to_wib(r["created_at"]) < end:
                    continue
                ids[r["id"]] = r["author_key"]
    authors.update(ids.values())
    return len(ids), len(authors)


def job_summary_text(job, root, task_ids=None, top=10):
    """Ringkasan teks untuk Telegram (pratinjau Excel di HP tidak selalu rapi)."""
    start = datetime.strptime(job["analysis_start_wib"], "%Y-%m-%d")
    end = datetime.strptime(job["analysis_end_wib"], "%Y-%m-%d") + timedelta(days=1)
    members = {m["key"]: m for m in job["members"]}
    tasks = selected_tasks(job, root, task_ids)
    lines = [f"🏆 Job {job['name']} ({job['analysis_start_wib']} – {job['analysis_end_wib']} WIB)"]
    ment = [(*_count_task(root, t, start, end), members[t["member"]]["panggilan"])
            for t in tasks if t["mode"] == "mention"]
    if ment:
        lines.append("Mention terbanyak (tweet · akun unik):")
        for i, (n, u, name) in enumerate(sorted(ment, reverse=True)[:top], 1):
            lines.append(f"{i}. {name} — {_id(n)} · {_id(u)}")
    tags = [(_count_task(root, t, start, end)[0], f"#{t.get('tag', '')} ({members[t['member']]['panggilan']})")
            for t in tasks if t["mode"] in ("hashtag", "tagline")]
    if tags:
        lines.append("Hashtag terbanyak (tweet):")
        for i, (n, name) in enumerate(sorted(tags, reverse=True)[:top], 1):
            lines.append(f"{i}. {name} — {_id(n)}")
    n_member = sum(1 for t in tasks if t["mode"] == "member")
    if n_member:
        lines.append(f"Metrik member: {n_member} member (lihat sheet Ringkasan).")
    return "\n".join(lines)[:4000]


def main():
    job = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
    print(build_job_workbook(job, Path(sys.argv[2]), Path(sys.argv[3])))
    print(job_summary_text(job, Path(sys.argv[2])))


if __name__ == "__main__":
    main()
