"""Tes offline job_worker/fetch_query/export_jobs/stress_test dengan respons GraphQL X tiruan.

Tidak ada request ke X/Telegram (Telegram diarahkan ke token palsu & gagal diam-diam).
Pakai: python tests/test_jobs_offline.py   (butuh twitter-cli, openpyxl)
"""
import hashlib
import json
import os
import random
import re
import shutil
import sys
import tempfile
import time
from datetime import datetime, timedelta, timezone
from pathlib import Path

BASE = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE))
TMP = Path(tempfile.mkdtemp(prefix="jobtest_"))
os.environ.update({"DATA_DIR": str(TMP / "data"), "TWITTER_AUTH_TOKEN": "x", "TWITTER_CT0": "y",
                   "TELEGRAM_BOT_TOKEN": "dummy", "TELEGRAM_CHAT_ID": "0", "EXIT_WHEN_DONE": "1",
                   "HASH_SALT": "garam-tes", "ALLOW_DIRECT": "1"})
os.environ.pop("TWITTER_PROXY", None)
os.environ.pop("REQUIRE_PROXY", None)

time.sleep = lambda s: None  # semua jeda/pacing dilewati
import fetch_v2  # noqa: E402
import worker as w  # noqa: E402
import job_worker  # noqa: E402
import stress_test  # noqa: E402
from xraw import parse_raw_any  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

WIB = timezone(timedelta(hours=7))
FAIL = []


def check(cond, msg):
    print(("OK  " if cond else "GAGAL ") + msg)
    if not cond:
        FAIL.append(msg)


# ── Data tiruan ───────────────────────────────────────────────────────
def user(uid, handle, followers=100):
    return {"rest_id": uid, "core": {"screen_name": handle, "created_at": "Mon Jan 01 00:00:00 +0000 2018"},
            "legacy": {"followers_count": followers}}


def tweet(tid, dt, author, text, mentions=(), tags=(), reply_to=None, likes=10, views=100):
    lg = {"id_str": tid, "created_at": dt.astimezone(timezone.utc).strftime("%a %b %d %H:%M:%S +0000 %Y"),
          "full_text": text, "favorite_count": likes, "retweet_count": 1, "reply_count": 0, "quote_count": 0,
          "bookmark_count": 0, "lang": "in",
          "entities": {"user_mentions": [{"screen_name": m} for m in mentions],
                       "hashtags": [{"text": t} for t in tags]}}
    if reply_to:
        lg["in_reply_to_status_id_str"] = "1"
        lg["in_reply_to_screen_name"] = reply_to
    return {"__typename": "Tweet", "legacy": lg, "core": {"user_results": {"result": author}},
            "views": {"count": str(views)}}


random.seed(7)
FENI = user("1", "F_FeniJKT48", 500000)
FANBASE = user("2", "Fenidelity_INA", 20000)
FANS = [user(str(100 + i), f"fan{i}", random.randint(10, 5000)) for i in range(30)]
T0 = datetime(2023, 12, 17, tzinfo=WIB)
CORPUS = []  # semua tweet tiruan sepanjang 4 hari pertama periode
for i in range(400):
    dt = T0 + timedelta(minutes=random.randint(0, 4 * 24 * 60 - 1))
    kind = random.random()
    if kind < 0.1:
        CORPUS.append(tweet(f"9{i:05d}", dt, FENI, f"halo semua {i}", tags=["CHANCE"] if i % 3 == 0 else []))
    elif kind < 0.25:
        CORPUS.append(tweet(f"9{i:05d}", dt, FANBASE, f"ayo dukung @F_FeniJKT48 #CHANCE {i}",
                            mentions=["F_FeniJKT48"], tags=["CHANCE"]))
    elif kind < 0.9:
        CORPUS.append(tweet(f"9{i:05d}", dt, random.choice(FANS), f"fans teks rahasia {i} @F_FeniJKT48",
                            mentions=["F_FeniJKT48"], tags=["CHANCE"] if i % 4 == 0 else [],
                            likes=random.randint(0, 500), views=random.randint(10, 20000)))
    else:  # reply ke member: harus tersaring -filter:replies
        CORPUS.append(tweet(f"9{i:05d}", dt, random.choice(FANS), f"@F_FeniJKT48 balasan {i}",
                            mentions=["F_FeniJKT48"], reply_to="F_FeniJKT48"))

# retweet polos fans atas tweet Feni: search 'Latest' memang mengembalikan RT, query memakai -filter:nativeretweets
FENI_TWEETS = [t for t in CORPUS if t["core"]["user_results"]["result"]["core"]["screen_name"] == "F_FeniJKT48"]
for i, orig in enumerate(FENI_TWEETS[:15]):
    rt = tweet(f"8{i:05d}", T0 + timedelta(hours=10 + i), random.choice(FANS), "RT @F_FeniJKT48: ...",
               mentions=["F_FeniJKT48"], tags=["CHANCE"], views=999999)
    rt["legacy"]["retweeted_status_result"] = {"result": orig}
    CORPUS.append(rt)
# fans mengutip fans lain & menyebut fans lain: handle fans lain tidak boleh tersimpan
q = tweet("700001", T0 + timedelta(hours=5), FANS[0], "kutip @F_FeniJKT48 @fan9", mentions=["F_FeniJKT48", "fan9"])
q["legacy"]["is_quote_status"] = True
q["quoted_status_result"] = {"result": tweet("700000", T0, FANS[9], "asli fan9")}
CORPUS.append(q)

DEPTH_CAP = int(os.environ.get("FAKE_DEPTH_CAP", "100000"))
CALLS = {"n": 0}


import threading  # noqa: E402
QUERY_LOG = []  # (nama thread, query) untuk memastikan dua worker tidak menarik query yang sama


def fake_search_page(client, query, cursor, on_rate_limit=None, on_transient=None, product="Latest"):
    """Tiru SearchTimeline: filter query sederhana, urut terbaru, 20 per halaman, cursor offset."""
    CALLS["n"] += 1
    QUERY_LOG.append((threading.current_thread().name, query))
    if os.environ.get("FAKE_LATENCY"):
        threading.Event().wait(float(os.environ["FAKE_LATENCY"]))  # supaya dua worker benar-benar bergantian
    m_from = re.findall(r"(?<!-)from:(\w+)", query)
    m_not_from = re.findall(r"-from:(\w+)", query)
    mentions = re.findall(r"@(\w+)", re.sub(r"-from:\w+", "", query))
    tag = re.search(r"#(\w+)", query)
    if "since_time" in query:
        lo = datetime.fromtimestamp(int(re.search(r"since_time:(\d+)", query).group(1)), timezone.utc)
        hi = datetime.fromtimestamp(int(re.search(r"until_time:(\d+)", query).group(1)), timezone.utc)
    else:
        lo = datetime.strptime(re.search(r"since:(\S+)", query).group(1), "%Y-%m-%d").replace(tzinfo=timezone.utc)
        hi = datetime.strptime(re.search(r"until:(\S+)", query).group(1), "%Y-%m-%d").replace(tzinfo=timezone.utc)
    out = []
    for t in CORPUS:
        lg, h = t["legacy"], t["core"]["user_results"]["result"]["core"]["screen_name"]
        dt = datetime.strptime(lg["created_at"], "%a %b %d %H:%M:%S %z %Y")
        if not lo <= dt < hi:
            continue
        if m_from and h.lower() not in {x.lower() for x in m_from}:
            continue
        if h.lower() in {x.lower() for x in m_not_from}:
            continue
        if "-filter:replies" in query and lg.get("in_reply_to_status_id_str"):
            continue
        if "-filter:nativeretweets" in query and lg.get("retweeted_status_result"):
            continue
        ents = lg["entities"]
        if tag:
            if tag.group(1).lower() not in {x["text"].lower() for x in ents["hashtags"]}:
                continue
        elif mentions and not m_from:
            if not {x.lower() for x in mentions} & {x["screen_name"].lower() for x in ents["user_mentions"]}:
                continue
        out.append(t)
    out.sort(key=lambda t: t["legacy"]["id_str"], reverse=True)
    out = out[:DEPTH_CAP]
    off = int(cursor or 0)
    page = out[off:off + 20]
    entries = [{"entryId": f"tweet-{t['legacy']['id_str']}",
                "content": {"itemContent": {"tweet_results": {"result": t}}}} for t in page]
    if off + 20 < len(out):
        entries.append({"entryId": "cursor-bottom-x", "content": {"value": str(off + 20)}})
    return {"data": {"search_by_raw_query": {"search_timeline": {"timeline": {"instructions": [
        {"type": "TimelineAddEntries", "entries": entries}]}}}}}


class FakeProfile:
    def __init__(self, h):
        self.id, self.screen_name, self.followers_count = "1", h, 500000


fetch_v2.search_page = fake_search_page
import twitter_cli.client as tc  # noqa: E402
tc.TwitterClient.fetch_user = lambda self, h: FakeProfile(h)
w.choose_proxy = lambda: ("", "tes-offline")
SENT = []


def _no_tg(method, fields, file_path=None):
    SENT.append((method, fields.get("text") or fields.get("caption"), file_path))
    raise RuntimeError("telegram dimatikan di tes")


w._tg = _no_tg


# ── 1. parser ─────────────────────────────────────────────────────────
classify_author = job_worker.make_classifier(job_worker.load_job())
hasher = lambda uid: "u_" + hashlib.sha256(f"s:{uid}".encode()).hexdigest()[:16]
fan_tweet = next(t for t in CORPUS if t["core"]["user_results"]["result"]["core"]["screen_name"].startswith("fan")
                 and not t["legacy"].get("in_reply_to_status_id_str"))
r = parse_raw_any(fan_tweet, classify_author, hasher)
check(r["author_type"] == "lain" and r["author"] == "" and r["text"] == "" and r["author_key"].startswith("u_"),
      "fans individu dianonimkan (tanpa handle & teks)")
fb = next(t for t in CORPUS if t["core"]["user_results"]["result"]["core"]["screen_name"] == "Fenidelity_INA")
r = parse_raw_any(fb, classify_author, hasher)
check(r["author_type"] == "fanbase" and r["author"] == "Fenidelity_INA" and r["text"], "fanbase tetap terbaca")
check(r["mentions"] == ["f_fenijkt48"] and r["hashtags"] == ["chance"], "mention & hashtag dari entities")

# ── 2. job_worker end-to-end (member + mention + hashtag Feni), periode dipersempit ke 4 hari ──
job = job_worker.load_job()
job["analysis_start_wib"], job["analysis_end_wib"] = "2023-12-17", "2023-12-20"
job["fetch_start_utc"], job["fetch_end_utc"] = "2023-12-16", "2023-12-21"
job_path = TMP / "job.json"
job_path.write_text(json.dumps(job))
os.environ.update({"JOB_FILE": str(job_path), "TASKS": "member:feni,mention:feni,hashtag:feni:chance",
                   "RUN_ID": "tes"})
job_worker.main()
root = TMP / "data" / "tes"
state_files = list(root.glob("state_antrean_*_*.json"))
check(len(state_files) == 1, "state mode antrean memakai sidik set task + nama worker")
state = json.loads(state_files[0].read_text())
check(state["finished"] == ["member:feni", "mention:feni", "hashtag:feni:chance"], "semua task selesai")
check(state.get("built", "").endswith(".xlsx"), "Excel dibuat")

expect_mention = {t["legacy"]["id_str"] for t in CORPUS
                  if "F_FeniJKT48" in [m["screen_name"] for m in t["legacy"]["entities"]["user_mentions"]]
                  and not t["legacy"].get("in_reply_to_status_id_str") and not t["legacy"].get("retweeted_status_result")
                  and t["core"]["user_results"]["result"]["core"]["screen_name"] != "F_FeniJKT48"}
got = set(fetch_v2.read_jsonl(root / "mention" / "mention__feni.jsonl"))
check(got == expect_mention, f"mention lengkap & tanpa reply ({len(got)}/{len(expect_mention)})")
expect_tag = {t["legacy"]["id_str"] for t in CORPUS if "CHANCE" in [h["text"] for h in t["legacy"]["entities"]["hashtags"]]
              and not t["legacy"].get("retweeted_status_result")}
got_tag = set(fetch_v2.read_jsonl(root / "hashtag" / "hashtag__feni__chance.jsonl"))
check(got_tag == expect_tag, f"hashtag lengkap ({len(got_tag)}/{len(expect_tag)})")
member_store = json.loads((root / "member" / "F_FeniJKT48_v2.json").read_text())
expect_member = {t["legacy"]["id_str"] for t in CORPUS
                 if t["core"]["user_results"]["result"]["core"]["screen_name"] == "F_FeniJKT48"}
check(set(member_store["tweets"]) == expect_member, "metrik member lengkap")
raw_text = (root / "mention" / "mention__feni.jsonl").read_text()
check("fans teks rahasia" not in raw_text, "teks fans individu tidak pernah ditulis ke disk")

wb = load_workbook(state["built"], read_only=True)
check(set(wb.sheetnames) >= {"Ringkasan", "Hashtag", "Harian", "Top Mention", "Top Hashtag", "Catatan"},
      f"sheet Excel lengkap {wb.sheetnames}")
rows = list(wb["Harian"].iter_rows(values_only=True))
check(rows[0][0] == "Tanggal (WIB)" and len(rows) == 1 + 4, "sheet Harian 4 hari (WIB)")
total_daily = sum(v for r in rows[1:] for v in r[1:] if isinstance(v, int))
check(total_daily == len(expect_mention) + len(expect_tag), "jumlah harian = mention + hashtag")
csvs = list((root / "exports").glob("mention_*.csv.gz"))
check(len(csvs) == 1, "data mentah CSV.gz per ekspor ditulis")
import gzip, csv as _csv  # noqa: E401,E402
csv_rows = list(_csv.DictReader(gzip.open(csvs[0], "rt", encoding="utf-8")))
fan_rows = [r for r in csv_rows if r["author_type"] == "lain"]
check(fan_rows and all(r["id_tweet"].startswith("t_") and not r["author"] and not r["text"]
                       and not r["author_followers"] for r in fan_rows), "CSV: id/handle/teks/follower fans dianonimkan")
raw_all = "".join(p.read_text() for p in (root / "mention").glob("*.jsonl"))
check('"fan9"' not in raw_all and "fan9" not in raw_all, "handle fans lain (quote/mention) tidak tersimpan")
check("999999" not in raw_all, "retweet polos tidak tersimpan sebagai mention")
top = list(wb["Top Mention"].iter_rows(values_only=True))[1:]
check(top and all((r[10] in ("", None)) and (r[6] in ("", None)) for r in top if r[4] == "lain"),
      "Top Mention: fans anonim tanpa link & follower")
ring = list(wb["Ringkasan"].iter_rows(values_only=True))
feni = next(r for r in ring[2:] if len(r) > 1 and r[1] == "Feni")
head = ring[1]
check(feni[head.index("Hashtag · tweet")] == len(expect_tag), "Ringkasan: tweet hashtag Feni tanpa dobel")
others = [r for r in ring[2:] if len(r) > 1 and r[1] != "Feni"]
check(all(r[head.index("Mention · tweet")] == "–" for r in others), "task yang tidak dijalankan tampil '–', bukan 0")
check(not [x for x in SENT if x[0] == "sendDocument" and x[2] is None], "tidak ada lalu lintas Telegram sungguhan")

# ── 3. restart: tidak menarik ulang ──────────────────────────────────
calls_before = CALLS["n"]
job_worker.main()
check(CALLS["n"] == calls_before, "restart setelah selesai tidak memanggil X lagi")

# ── 4. resume di tengah jendela: proses mati lalu dilanjutkan, tanpa duplikat ──
crash_root = TMP / "crash"
calls = {"n": 0}
orig = fetch_v2.search_page


def flaky(*a, **k):
    calls["n"] += 1
    if calls["n"] == 3:
        raise KeyboardInterrupt("mati di tengah")
    return orig(*a, **k)


fetch_v2.search_page = flaky
start = datetime(2023, 12, 17, tzinfo=WIB)
end = start + timedelta(days=4)
parse = lambda x: parse_raw_any(x, classify_author, hasher)
try:
    fetch_v2.fetch_query(None, "x", "(@F_FeniJKT48) -from:F_FeniJKT48 -filter:replies -filter:nativeretweets", parse, start, end, crash_root,
                         timedelta(hours=24), unit="time", pace=lambda: None, log=lambda s: None)
except KeyboardInterrupt:
    pass
fetch_v2.search_page = orig
fetch_v2.fetch_query(None, "x", "(@F_FeniJKT48) -from:F_FeniJKT48 -filter:replies -filter:nativeretweets", parse, start, end, crash_root,
                     timedelta(hours=24), unit="time", pace=lambda: None, log=lambda s: None)
got = set(fetch_v2.read_jsonl(crash_root / "x.jsonl"))
check(got == expect_mention, "resume setelah mati di tengah jendela: hasil tetap lengkap & unik")

# ── 5. stress test: batas kedalaman tiruan 30 tweet/query → jendela 24 jam kehilangan data ──
fetch_v2.search_page = lambda *a, **k: (globals().__setitem__("DEPTH_CAP", 30), fake_search_page(*a, **k))[1]
os.environ.update({"STRESS_TASKS": "mention:feni", "STRESS_START": "2023-12-17", "STRESS_DAYS": "1",
                   "STRESS_HOURS": "24,1", "RUN_ID": "stres"})
stress_test.main()
rep = json.loads((TMP / "data" / "stress" / "stres" / "report.json").read_text())
check(any("20231217_1d" in p.name for p in (TMP / "data" / "stress" / "stres").glob("*.jsonl")),
      "label uji batas memuat periode")
c24 = next(r for r in rep if r["jendela_jam"] == 24)
c1 = next(r for r in rep if r["jendela_jam"] == 1)
check(c24["tweet_unik"] < c1["tweet_unik"] and c24["cakupan_vs_gabungan"] < 1,
      f"uji batas mendeteksi kehilangan di jendela 24 jam ({c24['tweet_unik']} vs {c1['tweet_unik']})")

# ── 6. error parsial X: data + errors tetap dipakai; code 88 = rate limit; errors tanpa data = gagal ──
class Resp:
    def __init__(self, body, code=200):
        self.text, self.status_code = json.dumps(body), code


class Sess:
    def __init__(self, body, code=200):
        self.r = Resp(body, code)

    def post(self, *a, **k):
        return self.r


class Cl:
    def _build_headers(self, url="", method="GET"):
        return {}


fetch_v2.search_page = orig  # kembali ke search_page asli (memakai _post_tolerant)
from twitter_cli import client as _tc  # noqa: E402
from twitter_cli.exceptions import TwitterAPIError  # noqa: E402
_tc._get_cffi_session = lambda: Sess({"data": {"x": 1}, "errors": [{"code": 37, "message":
                                      "Authorization: Denied by access control: Missing TwitterUserNotSuspended"}]})
check(fetch_v2._post_tolerant(Cl(), "SearchTimeline", {}) == {"data": {"x": 1}, "errors": [{"code": 37, "message":
      "Authorization: Denied by access control: Missing TwitterUserNotSuspended"}]}, "error parsial + data -> data dipakai")
_tc._get_cffi_session = lambda: Sess({"errors": [{"code": 88, "message": "Rate limit exceeded"}]})
try:
    fetch_v2._post_tolerant(Cl(), "SearchTimeline", {})
    check(False, "code 88 harus jadi rate limit")
except TwitterAPIError as e:
    check(fetch_v2.classify(e) == "rate", "code 88 -> rate limit")
check(fetch_v2.classify(TwitterAPIError(0, "Twitter API returned errors: Missing TwitterUserNotSuspended")) != "account",
      "'TwitterUserNotSuspended' tidak dianggap masalah akun")
check(fetch_v2.classify(TwitterAPIError(0, "Your account is suspended and is not permitted")) == "account",
      "akun suspend tetap terdeteksi")

# ── 7. salt: env berbeda dari salt tersimpan -> berhenti; shard >1 tanpa salt -> berhenti ──
sroot = TMP / "salt"
sroot.mkdir()
os.environ["HASH_SALT"] = "a"
os.environ["SHARD"] = "1/1"
check(job_worker.hash_salt(sroot) == "a", "salt env ditulis ke file")
os.environ["HASH_SALT"] = "b"
try:
    job_worker.hash_salt(sroot)
    check(False, "salt berbeda harus ditolak")
except w.Stop:
    check(True, "salt berbeda dari yang tersimpan ditolak")
os.environ["HASH_SALT"] = ""
os.environ["SHARD"] = "1/2"
try:
    job_worker.hash_salt(TMP / "salt2" if (TMP / "salt2").mkdir() is None else None)
    check(False, "shard >1 tanpa salt harus ditolak")
except w.Stop:
    check(True, "shard >1 tanpa HASH_SALT ditolak")
os.environ["SHARD"] = "1/1"

# ── 8. dua worker paralel (2 akun) di antrean bersama: tiap task tepat sekali ──
from taskqueue import TaskQueue  # noqa: E402
fetch_v2.search_page = fake_search_page
globals()["DEPTH_CAP"] = 100000
QUERY_LOG.clear()
par_tasks = ["member:feni", "mention:feni", "hashtag:feni:chance", "mention:christy", "mention:indah",
             "mention:delynn", "mention:freya", "mention:muthe"]
os.environ.update({"TASKS": ",".join(par_tasks), "RUN_ID": "paralel", "SHARD": "", "HASH_SALT": "",
                   "QUEUE_POLL_S": "0.05", "FAKE_LATENCY": "0.01"})
orig_wid = job_worker.worker_id
job_worker.worker_id = lambda: threading.current_thread().name
errs = []


def run_worker():
    try:
        job_worker.main()
    except Exception as e:  # noqa: BLE001
        errs.append(repr(e))


ths = [threading.Thread(target=run_worker, name=n) for n in ("akun1", "akun2")]
for th in ths:
    th.start()
for th in ths:
    th.join(timeout=300)
job_worker.worker_id = orig_wid
os.environ.pop("FAKE_LATENCY")
proot = TMP / "data" / "paralel"
q = TaskQueue(proot, "pemeriksa")
check(not errs and all(not th.is_alive() for th in ths), f"kedua worker selesai tanpa error {errs}")
check(all(q.is_done(t) for t in par_tasks), "semua task paralel selesai")
by = {t: q.done_by(t) for t in par_tasks}
check(set(by.values()) == {"akun1", "akun2"}, f"kedua akun ikut bekerja {sorted(set(by.values()))}")
per_query_threads = {}
for th_name, qq in QUERY_LOG:
    per_query_threads.setdefault(qq, set()).add(th_name)
dup = [qq for qq, names in per_query_threads.items() if len(names) > 1]
check(not dup, f"tidak ada query yang ditarik dua akun ({len(dup)} dobel)")
xlsx = list(proot.glob("JKT48_paralel_*.xlsx"))
check(len(xlsx) == 1, f"Excel akhir dibuat tepat sekali ({len(xlsx)})")
got = set(fetch_v2.read_jsonl(proot / "mention" / "mention__feni.jsonl"))
check(got == expect_mention, "hasil mention Feni tetap lengkap di mode paralel")
check((proot / ".hash_salt").exists(), "salt dibuat sekali & dibagi lewat file")

# ── 9. klaim basi (worker mati) diambil alih; klaim hidup tidak ──
qroot = TMP / "q"
qa, qb = TaskQueue(qroot, "mati", stale_s=600), TaskQueue(qroot, "hidup", stale_s=600)
check(qa.claim("mention:feni"), "akun A mengklaim task")
check(not qb.claim("mention:feni"), "akun B tidak bisa mengambil task yang masih diklaim A")
claim_file = qroot / "queue" / "claims" / "mention__feni.json"
old_t = time.time() - 3600
os.utime(claim_file, (old_t, old_t))
check(qb.claim("mention:feni"), "klaim A yang basi (A mati) diambil alih B")
check(not qa.claim("mention:feni") or qa.claim_info("mention:feni")["worker"] == "hidup",
      "setelah diambil alih, A tidak merebut balik klaim yang masih hidup")
qb.finish("mention:feni")
check(qb.is_done("mention:feni") and not qa.claim("mention:feni"), "task selesai tidak diklaim lagi")
check(qa.take_export() and not qb.take_export(), "hanya satu worker yang membuat Excel akhir")

# ── 10. pecah otomatis: jendela besar yang padat ditarik ulang per hari → hasil lengkap ──
globals()["DEPTH_CAP"] = 50  # batas kedalaman tiruan: satu query hanya memberi 50 tweet teratas
start = datetime(2023, 12, 17, tzinfo=WIB)
qtag = "#CHANCE -filter:nativeretweets"
st_no = fetch_v2.fetch_query(None, "tanpa_pecah", qtag, parse, start, start + timedelta(days=4), TMP / "split",
                             timedelta(hours=96), unit="time", pace=lambda: None, log=lambda s: None)
st_yes = fetch_v2.fetch_query(None, "pecah", qtag, parse, start, start + timedelta(days=4), TMP / "split",
                              timedelta(hours=96), unit="time", pace=lambda: None, log=lambda s: None,
                              split_window=timedelta(hours=24), split_threshold=40)
n_no = len(fetch_v2.read_jsonl(TMP / "split" / "tanpa_pecah.jsonl"))
n_yes = len(fetch_v2.read_jsonl(TMP / "split" / "pecah.jsonl"))
check(n_no < len(expect_tag) and n_yes == len(expect_tag),
      f"jendela padat dipecah otomatis: {n_no} -> {n_yes}/{len(expect_tag)} tweet")
check(sum(1 for x in st_yes["windows"].values() if x["sub"]) == 4, "4 sub-jendela harian tercatat")
globals()["DEPTH_CAP"] = 100000

shutil.rmtree(TMP, ignore_errors=True)
print("\nSEMUA LULUS" if not FAIL else f"\n{len(FAIL)} GAGAL: {FAIL}")
sys.exit(1 if FAIL else 0)
